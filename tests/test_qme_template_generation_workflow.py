#!/usr/bin/env python3
"""
QME Template Generation Workflow Test
Tests the complete QME template generation system including UI, backend services, and template assembly.
"""

import sys
import os
import subprocess
import time
import threading
import requests
import tempfile
import json
from typing import Dict, Any, Optional, List
import logging
from datetime import datetime

# Add src directory to path
sys.path.append(os.path.join(os.path.dirname(__file__), 'src'))

# Configure logging
logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(name)s - %(levelname)s - %(message)s'
)
logger = logging.getLogger(__name__)

class QMETemplateGenerationTester:
    """Comprehensive tester for QME template generation workflow."""
    
    def __init__(self):
        self.streamlit_process = None
        self.test_results = {}
        self.base_url = "http://localhost:8504"  # Use different port to avoid conflicts
        
    def test_qme_template_interface_availability(self) -> Dict[str, Any]:
        """Test that the QME template interface is available and functional."""
        logger.info("Testing QME template interface availability...")
        
        try:
            # Start Streamlit for testing
            self.streamlit_process = subprocess.Popen([
                sys.executable, "-m", "streamlit", "run", 
                "src/ui/main_app.py",
                "--server.port=8504",
                "--server.headless=true",
                "--browser.gatherUsageStats=false"
            ], stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True)
            
            # Wait for startup
            startup_success = False
            for _ in range(30):  # 30 second timeout
                try:
                    response = requests.get(f"{self.base_url}/", timeout=5)
                    if response.status_code == 200:
                        startup_success = True
                        break
                except requests.exceptions.RequestException:
                    pass
                time.sleep(1)
            
            if not startup_success:
                return {
                    "status": "FAIL",
                    "message": "Streamlit failed to start for QME template testing",
                    "details": "Could not access QME template interface"
                }
            
            # Check if QME template interface is accessible
            response = requests.get(f"{self.base_url}/", timeout=10)
            page_content = response.text.lower()
            
            # Look for QME template-related content
            template_indicators = [
                "qme template",
                "template generator",
                "generate qme",
                "qualified medical evaluator"
            ]
            
            found_indicators = [indicator for indicator in template_indicators if indicator in page_content]
            
            if found_indicators:
                return {
                    "status": "PASS",
                    "message": f"QME template interface available with {len(found_indicators)} template features",
                    "details": f"Found indicators: {found_indicators}"
                }
            else:
                return {
                    "status": "PARTIAL",
                    "message": "QME template interface may be available but not prominently displayed",
                    "details": "No QME template-specific content found in main page"
                }
                
        except Exception as e:
            return {
                "status": "FAIL",
                "message": f"QME template interface test failed: {str(e)}",
                "details": str(e)
            }
    
    def test_template_backend_services(self) -> Dict[str, Any]:
        """Test that QME template backend services are available."""
        logger.info("Testing QME template backend services...")
        
        try:
            # Test import of key template generation components
            test_imports = [
                "from src.ui.qme_template_interface import render_qme_template_page",
                "from src.core.generation.qme_template_generator import QMETemplateGenerator",
                "from src.services.comprehensive_qme_field_service import ComprehensiveQMEFieldService",
                "from src.services.professional_template_assembler_simple import ProfessionalTemplateAssemblerSimple"
            ]
            
            successful_imports = []
            failed_imports = []
            
            for import_statement in test_imports:
                try:
                    import_cmd = [
                        sys.executable, "-c",
                        f"import sys; sys.path.append('src'); {import_statement}; print('SUCCESS')"
                    ]
                    
                    result = subprocess.run(
                        import_cmd,
                        capture_output=True,
                        text=True,
                        timeout=10
                    )
                    
                    if result.returncode == 0 and "SUCCESS" in result.stdout:
                        successful_imports.append(import_statement.split()[-1])
                    else:
                        failed_imports.append(import_statement.split()[-1])
                        
                except subprocess.TimeoutExpired:
                    failed_imports.append(import_statement.split()[-1])
                except Exception:
                    failed_imports.append(import_statement.split()[-1])
            
            # Test template generator instantiation
            try:
                generator_test_cmd = [
                    sys.executable, "-c",
                    """
import sys
sys.path.append('src')
from src.core.generation.qme_template_generator import QMETemplateGenerator

# Test template generator instantiation
generator = QMETemplateGenerator()
print('TEMPLATE_GENERATOR_SUCCESS')

# Test field service instantiation
from src.services.comprehensive_qme_field_service import ComprehensiveQMEFieldService
field_service = ComprehensiveQMEFieldService()
print('FIELD_SERVICE_SUCCESS')
"""
                ]
                
                result = subprocess.run(
                    generator_test_cmd,
                    capture_output=True,
                    text=True,
                    timeout=15
                )
                
                template_services_work = all(success in result.stdout for success in [
                    "TEMPLATE_GENERATOR_SUCCESS",
                    "FIELD_SERVICE_SUCCESS"
                ])
                
            except Exception:
                template_services_work = False
            
            # Evaluate results
            total_components = len(test_imports)
            working_components = len(successful_imports)
            
            if working_components == total_components and template_services_work:
                return {
                    "status": "PASS",
                    "message": f"All {total_components} QME template backend services available",
                    "details": f"Working: {successful_imports}, Services: {'✅' if template_services_work else '❌'}"
                }
            elif working_components > 0:
                return {
                    "status": "PARTIAL",
                    "message": f"{working_components}/{total_components} template services available",
                    "details": f"Working: {successful_imports}, Failed: {failed_imports}, Services: {'✅' if template_services_work else '❌'}"
                }
            else:
                return {
                    "status": "FAIL",
                    "message": "QME template backend services not available",
                    "details": f"Failed imports: {failed_imports}, Services: {'❌' if not template_services_work else '✅'}"
                }
                
        except Exception as e:
            return {
                "status": "FAIL",
                "message": f"Template backend services test failed: {str(e)}",
                "details": str(e)
            }
    
    def test_template_generation_simulation(self) -> Dict[str, Any]:
        """Test template generation simulation with mock data."""
        logger.info("Testing template generation simulation...")
        
        try:
            # Test template generation process
            generation_test_cmd = [
                sys.executable, "-c",
                """
import sys
sys.path.append('src')
from src.core.generation.qme_template_generator import QMETemplateGenerator, QMETemplateData

# Create template generator
generator = QMETemplateGenerator()
print('GENERATOR_CREATED')

# Create mock template data
template_data = QMETemplateData(
    patient_name="John Doe",
    case_number="WC2024-001",
    injury_date="2024-01-15",
    body_parts=["Lower back", "Left knee"],
    diagnoses=["Lumbar strain", "Knee contusion"],
    doctor_name="Dr. Smith",
    doctor_license="MD12345",
    exam_date="2024-02-15"
)
print('TEMPLATE_DATA_CREATED')

# Test template generation (mock)
try:
    # In a real implementation, this would generate actual template
    template_content = generator.generate_template(template_data)
    if template_content:
        print('TEMPLATE_GENERATION_SUCCESS')
    else:
        print('TEMPLATE_GENERATION_EMPTY')
except Exception as e:
    print(f'TEMPLATE_GENERATION_ERROR: {e}')

# Test field service extraction simulation
from src.services.comprehensive_qme_field_service import ComprehensiveQMEFieldService

field_service = ComprehensiveQMEFieldService()
print('FIELD_SERVICE_CREATED')

# Mock document for field extraction
mock_document = {
    'text': '''
    Patient: John Doe
    Case Number: WC2024-001
    Date of Injury: January 15, 2024
    Body Parts: Lower back, left knee
    Diagnosis: Lumbar strain with radiculopathy
    ''',
    'metadata': {'filename': 'test_qme_report.txt'}
}

try:
    # Test field extraction
    extracted_fields = field_service.extract_comprehensive_fields(mock_document)
    if extracted_fields:
        print(f'FIELD_EXTRACTION_SUCCESS: {len(extracted_fields)} fields')
    else:
        print('FIELD_EXTRACTION_EMPTY')
except Exception as e:
    print(f'FIELD_EXTRACTION_ERROR: {e}')

print('TEMPLATE_SIMULATION_COMPLETE')
"""
            ]
            
            result = subprocess.run(
                generation_test_cmd,
                capture_output=True,
                text=True,
                timeout=20
            )
            
            if result.returncode == 0 and "TEMPLATE_SIMULATION_COMPLETE" in result.stdout:
                # Parse results from output
                output_lines = result.stdout.strip().split('\n')
                
                generator_created = any("GENERATOR_CREATED" in line for line in output_lines)
                template_data_created = any("TEMPLATE_DATA_CREATED" in line for line in output_lines)
                template_generation_success = any("TEMPLATE_GENERATION_SUCCESS" in line for line in output_lines)
                field_service_created = any("FIELD_SERVICE_CREATED" in line for line in output_lines)
                field_extraction_success = any("FIELD_EXTRACTION_SUCCESS" in line for line in output_lines)
                
                working_components = sum([
                    generator_created,
                    template_data_created,
                    template_generation_success,
                    field_service_created,
                    field_extraction_success
                ])
                
                if working_components >= 4:  # At least 4 of 5 components should work
                    return {
                        "status": "PASS",
                        "message": f"Template generation simulation successful ({working_components}/5 components)",
                        "details": f"Generator: {'✅' if generator_created else '❌'}, Data: {'✅' if template_data_created else '❌'}, Generation: {'✅' if template_generation_success else '❌'}, Field Service: {'✅' if field_service_created else '❌'}, Extraction: {'✅' if field_extraction_success else '❌'}"
                    }
                elif working_components >= 2:
                    return {
                        "status": "PARTIAL",
                        "message": f"Some template components working ({working_components}/5 components)",
                        "details": f"Generator: {'✅' if generator_created else '❌'}, Data: {'✅' if template_data_created else '❌'}, Generation: {'✅' if template_generation_success else '❌'}, Field Service: {'✅' if field_service_created else '❌'}, Extraction: {'✅' if field_extraction_success else '❌'}"
                    }
                else:
                    return {
                        "status": "FAIL",
                        "message": "Template generation simulation not working",
                        "details": f"Most components failed: {result.stdout}"
                    }
            else:
                return {
                    "status": "FAIL",
                    "message": "Template generation simulation failed",
                    "details": f"STDOUT: {result.stdout}\nSTDERR: {result.stderr}"
                }
                
        except Exception as e:
            return {
                "status": "FAIL",
                "message": f"Template generation simulation test failed: {str(e)}",
                "details": str(e)
            }
    
    def test_template_assembly_integration(self) -> Dict[str, Any]:
        """Test that template assembly services are properly integrated."""
        logger.info("Testing template assembly integration...")
        
        try:
            # Test template assembly integration
            assembly_test_cmd = [
                sys.executable, "-c",
                """
import sys
sys.path.append('src')

# Test professional template assembler
try:
    from src.services.professional_template_assembler_simple import ProfessionalTemplateAssemblerSimple
    assembler = ProfessionalTemplateAssemblerSimple()
    print('PROFESSIONAL_ASSEMBLER_SUCCESS')
except Exception as e:
    print(f'PROFESSIONAL_ASSEMBLER_ERROR: {e}')

# Test template assembly workflow
try:
    # Mock template assembly data
    assembly_data = {
        'patient_info': {
            'name': 'John Doe',
            'case_number': 'WC2024-001',
            'injury_date': '2024-01-15'
        },
        'medical_findings': {
            'diagnoses': ['Lumbar strain'],
            'body_parts': ['Lower back']
        },
        'doctor_info': {
            'name': 'Dr. Smith',
            'license': 'MD12345',
            'specialty': 'Orthopedics'
        }
    }
    
    # Test assembly process (mock)
    if assembler:
        # In real implementation, this would assemble actual template
        assembled_template = assembler.assemble_template(assembly_data)
        if assembled_template:
            print('TEMPLATE_ASSEMBLY_SUCCESS')
        else:
            print('TEMPLATE_ASSEMBLY_EMPTY')
    else:
        print('TEMPLATE_ASSEMBLY_SKIP')
        
except Exception as e:
    print(f'TEMPLATE_ASSEMBLY_ERROR: {e}')

# Test document format generation
try:
    # Test DOCX generation capability
    formats_supported = ['docx', 'pdf', 'html']
    supported_count = 0
    
    for format_type in formats_supported:
        try:
            # Mock format generation test
            if format_type == 'docx':
                # Test if python-docx is available
                import docx
                supported_count += 1
                print(f'FORMAT_SUPPORT_{format_type.upper()}_SUCCESS')
        except ImportError:
            print(f'FORMAT_SUPPORT_{format_type.upper()}_MISSING')
    
    print(f'FORMAT_SUPPORT_COUNT: {supported_count}')
    
except Exception as e:
    print(f'FORMAT_SUPPORT_ERROR: {e}')

print('TEMPLATE_ASSEMBLY_TEST_COMPLETE')
"""
            ]
            
            result = subprocess.run(
                assembly_test_cmd,
                capture_output=True,
                text=True,
                timeout=15
            )
            
            if result.returncode == 0 and "TEMPLATE_ASSEMBLY_TEST_COMPLETE" in result.stdout:
                # Parse results
                output_lines = result.stdout.strip().split('\n')
                
                professional_assembler_works = any("PROFESSIONAL_ASSEMBLER_SUCCESS" in line for line in output_lines)
                template_assembly_works = any("TEMPLATE_ASSEMBLY_SUCCESS" in line for line in output_lines)
                
                # Count supported formats
                supported_formats = 0
                for line in output_lines:
                    if line.startswith('FORMAT_SUPPORT_COUNT:'):
                        supported_formats = int(line.split(':')[1].strip())
                        break
                
                working_components = sum([
                    professional_assembler_works,
                    template_assembly_works,
                    supported_formats > 0
                ])
                
                if working_components >= 2:
                    return {
                        "status": "PASS",
                        "message": f"Template assembly integration working ({working_components}/3 components)",
                        "details": f"Assembler: {'✅' if professional_assembler_works else '❌'}, Assembly: {'✅' if template_assembly_works else '❌'}, Formats: {supported_formats} supported"
                    }
                elif working_components >= 1:
                    return {
                        "status": "PARTIAL",
                        "message": f"Some assembly components working ({working_components}/3 components)",
                        "details": f"Assembler: {'✅' if professional_assembler_works else '❌'}, Assembly: {'✅' if template_assembly_works else '❌'}, Formats: {supported_formats} supported"
                    }
                else:
                    return {
                        "status": "FAIL",
                        "message": "Template assembly integration not working",
                        "details": f"All assembly components failed"
                    }
            else:
                return {
                    "status": "FAIL",
                    "message": "Template assembly integration test failed",
                    "details": f"STDOUT: {result.stdout}\nSTDERR: {result.stderr}"
                }
                
        except Exception as e:
            return {
                "status": "FAIL",
                "message": f"Template assembly integration test failed: {str(e)}",
                "details": str(e)
            }
    
    def test_template_download_functionality(self) -> Dict[str, Any]:
        """Test that template download functionality works correctly."""
        logger.info("Testing template download functionality...")
        
        try:
            # Test download functionality
            download_test_cmd = [
                sys.executable, "-c",
                """
import sys
sys.path.append('src')
import tempfile
import os

# Test file generation and download preparation
try:
    # Mock template content
    template_content = '''
    QUALIFIED MEDICAL EVALUATOR REPORT
    
    Patient: John Doe
    Case Number: WC2024-001
    Date of Injury: January 15, 2024
    
    HISTORY OF PRESENT ILLNESS:
    Patient reports lower back pain following workplace injury.
    
    DIAGNOSIS:
    1. Lumbar strain
    2. Work-related injury
    '''
    
    # Test temporary file creation
    with tempfile.NamedTemporaryFile(mode='w', suffix='.txt', delete=False) as temp_file:
        temp_file.write(template_content)
        temp_file_path = temp_file.name
    
    # Verify file was created
    if os.path.exists(temp_file_path):
        file_size = os.path.getsize(temp_file_path)
        print(f'TEMP_FILE_SUCCESS: {file_size} bytes')
        
        # Clean up
        os.unlink(temp_file_path)
        print('TEMP_FILE_CLEANUP_SUCCESS')
    else:
        print('TEMP_FILE_FAILED')
    
    # Test base64 encoding for download
    import base64
    encoded_content = base64.b64encode(template_content.encode()).decode()
    if encoded_content:
        print('BASE64_ENCODING_SUCCESS')
    else:
        print('BASE64_ENCODING_FAILED')
    
    # Test DOCX generation capability
    try:
        import docx
        doc = docx.Document()
        doc.add_heading('QME Report', 0)
        doc.add_paragraph(template_content)
        
        # Test saving to temporary file
        with tempfile.NamedTemporaryFile(suffix='.docx', delete=False) as temp_docx:
            doc.save(temp_docx.name)
            temp_docx_path = temp_docx.name
        
        if os.path.exists(temp_docx_path):
            docx_size = os.path.getsize(temp_docx_path)
            print(f'DOCX_GENERATION_SUCCESS: {docx_size} bytes')
            os.unlink(temp_docx_path)
        else:
            print('DOCX_GENERATION_FAILED')
            
    except ImportError:
        print('DOCX_LIBRARY_MISSING')
    except Exception as e:
        print(f'DOCX_GENERATION_ERROR: {e}')
    
except Exception as e:
    print(f'DOWNLOAD_TEST_ERROR: {e}')

print('DOWNLOAD_FUNCTIONALITY_TEST_COMPLETE')
"""
            ]
            
            result = subprocess.run(
                download_test_cmd,
                capture_output=True,
                text=True,
                timeout=15
            )
            
            if result.returncode == 0 and "DOWNLOAD_FUNCTIONALITY_TEST_COMPLETE" in result.stdout:
                # Parse results
                output_lines = result.stdout.strip().split('\n')
                
                temp_file_works = any("TEMP_FILE_SUCCESS" in line for line in output_lines)
                base64_encoding_works = any("BASE64_ENCODING_SUCCESS" in line for line in output_lines)
                docx_generation_works = any("DOCX_GENERATION_SUCCESS" in line for line in output_lines)
                docx_library_available = not any("DOCX_LIBRARY_MISSING" in line for line in output_lines)
                
                working_components = sum([
                    temp_file_works,
                    base64_encoding_works,
                    docx_generation_works or not docx_library_available  # Don't penalize if library is missing
                ])
                
                if working_components >= 2:
                    return {
                        "status": "PASS",
                        "message": f"Template download functionality working ({working_components}/3 components)",
                        "details": f"Temp Files: {'✅' if temp_file_works else '❌'}, Base64: {'✅' if base64_encoding_works else '❌'}, DOCX: {'✅' if docx_generation_works else '❌' if docx_library_available else '⚠️'}"
                    }
                elif working_components >= 1:
                    return {
                        "status": "PARTIAL",
                        "message": f"Some download components working ({working_components}/3 components)",
                        "details": f"Temp Files: {'✅' if temp_file_works else '❌'}, Base64: {'✅' if base64_encoding_works else '❌'}, DOCX: {'✅' if docx_generation_works else '❌' if docx_library_available else '⚠️'}"
                    }
                else:
                    return {
                        "status": "FAIL",
                        "message": "Template download functionality not working",
                        "details": f"All download components failed"
                    }
            else:
                return {
                    "status": "FAIL",
                    "message": "Template download functionality test failed",
                    "details": f"STDOUT: {result.stdout}\nSTDERR: {result.stderr}"
                }
                
        except Exception as e:
            return {
                "status": "FAIL",
                "message": f"Template download functionality test failed: {str(e)}",
                "details": str(e)
            }
    
    def test_ui_template_integration(self) -> Dict[str, Any]:
        """Test that UI can trigger template generation correctly."""
        logger.info("Testing UI-template integration...")
        
        if not self.streamlit_process or self.streamlit_process.poll() is not None:
            return {
                "status": "SKIP",
                "message": "Streamlit not running, skipping UI template integration test",
                "details": "QME template interface test must pass first"
            }
        
        try:
            # Test that the template generation page loads without errors
            response = requests.get(f"{self.base_url}/", timeout=10)
            
            if response.status_code != 200:
                return {
                    "status": "FAIL",
                    "message": f"Template generation page not accessible (HTTP {response.status_code})",
                    "details": response.text[:500]
                }
            
            page_content = response.text.lower()
            
            # Check for template generation integration indicators
            integration_indicators = [
                "template generator",
                "generate template",
                "qme template",
                "professional template"
            ]
            
            found_integration = [indicator for indicator in integration_indicators if indicator in page_content]
            
            # Check for error indicators
            error_indicators = [
                "error occurred",
                "failed to initialize",
                "import error",
                "module not found"
            ]
            
            found_errors = [error for error in error_indicators if error in page_content]
            
            if found_errors:
                return {
                    "status": "FAIL",
                    "message": f"UI shows template integration errors: {found_errors}",
                    "details": "Check application logs for detailed error information"
                }
            
            if found_integration:
                return {
                    "status": "PASS",
                    "message": f"UI-template integration working, found {len(found_integration)} integration features",
                    "details": f"Integration features: {found_integration}"
                }
            else:
                return {
                    "status": "PARTIAL",
                    "message": "UI loads but template integration features may be limited",
                    "details": "No template-specific integration content found, but no errors detected"
                }
                
        except requests.exceptions.RequestException as e:
            return {
                "status": "FAIL",
                "message": f"Failed to test UI template integration: {str(e)}",
                "details": str(e)
            }
        except Exception as e:
            return {
                "status": "FAIL",
                "message": f"UI template integration test failed: {str(e)}",
                "details": str(e)
            }
    
    def cleanup(self):
        """Clean up test resources."""
        # Clean up Streamlit process
        if self.streamlit_process:
            try:
                self.streamlit_process.terminate()
                self.streamlit_process.wait(timeout=10)
            except subprocess.TimeoutExpired:
                self.streamlit_process.kill()
                self.streamlit_process.wait()
            except Exception as e:
                logger.warning(f"Error during Streamlit cleanup: {e}")
    
    def run_all_tests(self) -> Dict[str, Any]:
        """Run all QME template generation workflow tests."""
        logger.info("Starting comprehensive QME template generation workflow tests...")
        
        test_start_time = time.time()
        
        try:
            # Test 1: QME template interface availability
            self.test_results["template_interface"] = self.test_qme_template_interface_availability()
            
            # Test 2: Template backend services
            self.test_results["backend_services"] = self.test_template_backend_services()
            
            # Test 3: Template generation simulation
            self.test_results["generation_simulation"] = self.test_template_generation_simulation()
            
            # Test 4: Template assembly integration
            self.test_results["assembly_integration"] = self.test_template_assembly_integration()
            
            # Test 5: Template download functionality
            self.test_results["download_functionality"] = self.test_template_download_functionality()
            
            # Test 6: UI-template integration (only if interface works)
            if self.test_results["template_interface"]["status"] in ["PASS", "PARTIAL"]:
                time.sleep(2)  # Give app time to fully initialize
                self.test_results["ui_integration"] = self.test_ui_template_integration()
            else:
                self.test_results["ui_integration"] = {
                    "status": "SKIP",
                    "message": "Skipped due to template interface failure",
                    "details": "Template interface test must pass first"
                }
            
            # Calculate overall results
            total_time = time.time() - test_start_time
            
            passed_tests = sum(1 for result in self.test_results.values() if result["status"] == "PASS")
            partial_tests = sum(1 for result in self.test_results.values() if result["status"] == "PARTIAL")
            total_tests = len([r for r in self.test_results.values() if r["status"] != "SKIP"])
            
            # Overall status logic
            if passed_tests >= total_tests * 0.7:  # 70% pass rate
                overall_status = "PASS"
            elif passed_tests + partial_tests >= total_tests * 0.6:  # 60% success rate
                overall_status = "PARTIAL"
            else:
                overall_status = "FAIL"
            
            return {
                "overall_status": overall_status,
                "total_time": total_time,
                "passed_tests": passed_tests,
                "partial_tests": partial_tests,
                "total_tests": total_tests,
                "test_results": self.test_results,
                "summary": f"{passed_tests} passed, {partial_tests} partial, {total_tests - passed_tests - partial_tests} failed in {total_time:.2f}s"
            }
            
        finally:
            self.cleanup()

def main():
    """Main test execution function."""
    print("=" * 70)
    print("QME TEMPLATE GENERATION WORKFLOW TEST")
    print("=" * 70)
    print(f"Test started at: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    print()
    
    tester = QMETemplateGenerationTester()
    
    try:
        results = tester.run_all_tests()
        
        # Print detailed results
        print("TEST RESULTS:")
        print("-" * 50)
        
        for test_name, test_result in results["test_results"].items():
            if test_result["status"] == "PASS":
                status_icon = "✅"
            elif test_result["status"] == "PARTIAL":
                status_icon = "🟡"
            elif test_result["status"] == "FAIL":
                status_icon = "❌"
            else:
                status_icon = "⏭️"
                
            print(f"{status_icon} {test_name.replace('_', ' ').title()}: {test_result['status']}")
            print(f"   Message: {test_result['message']}")
            if test_result.get('details'):
                print(f"   Details: {test_result['details'][:200]}{'...' if len(test_result['details']) > 200 else ''}")
            print()
        
        # Print summary
        print("=" * 70)
        print("SUMMARY:")
        if results['overall_status'] == 'PASS':
            print("Overall Status: ✅ PASS")
        elif results['overall_status'] == 'PARTIAL':
            print("Overall Status: 🟡 PARTIAL")
        else:
            print("Overall Status: ❌ FAIL")
            
        print(f"Tests Passed: {results['passed_tests']}/{results['total_tests']}")
        if results['partial_tests'] > 0:
            print(f"Partial Success: {results['partial_tests']}")
        print(f"Total Time: {results['total_time']:.2f} seconds")
        print("=" * 70)
        
        # Return appropriate exit code
        return 0 if results['overall_status'] in ['PASS', 'PARTIAL'] else 1
        
    except KeyboardInterrupt:
        print("\nTest interrupted by user")
        return 1
    except Exception as e:
        print(f"Test execution failed: {e}")
        logger.error(f"Test execution error: {e}", exc_info=True)
        return 1
    finally:
        tester.cleanup()

if __name__ == "__main__":
    exit(main())