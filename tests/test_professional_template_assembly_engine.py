"""
Tests for Professional Template Assembly Engine.

Tests validate template assembly quality against existing successful templates
and ensure compliance with requirements 3.1, 3.2, and 3.3.
"""

import pytest
import tempfile
import os
from pathlib import Path
from datetime import datetime
from unittest.mock import Mock, patch, MagicMock

try:
    from src.services.professional_template_assembly_engine import (
        ProfessionalTemplateAssemblyEngine,
        TemplateConfig,
        FormattingRules,
        QualityRequirements,
        AssemblyResult,
        AssemblyStatus,
        FallbackStrategy,
        QualityReport,
        TemplateValidator,
        ProfessionalFormatter,
        ErrorRecoveryManager
    )
    from src.services.qme_template_generator import QMETemplateData, PatientInfo, MedicalFindings
    from src.models.knowledge_graph import Diagnosis, ImpairmentRating, Finding
    from src.services.qme_rules_engine import ValidationIssue, QualityScore, ValidationSeverity
    from src.models.extraction_models import ExtractionResult
except ImportError:
    from services.professional_template_assembly_engine import (
        ProfessionalTemplateAssemblyEngine,
        TemplateConfig,
        FormattingRules,
        QualityRequirements,
        AssemblyResult,
        AssemblyStatus,
        FallbackStrategy,
        QualityReport,
        TemplateValidator,
        ProfessionalFormatter,
        ErrorRecoveryManager
    )
    from services.qme_template_generator import QMETemplateData, PatientInfo, MedicalFindings
    from models.knowledge_graph import Diagnosis, ImpairmentRating, Finding
    from services.qme_rules_engine import ValidationIssue, QualityScore, ValidationSeverity
    from models.extraction_models import ExtractionResult


class TestProfessionalTemplateAssemblyEngine:
    """Test suite for Professional Template Assembly Engine."""
    
    @pytest.fixture
    def sample_template_data(self):
        """Create sample template data for testing."""
        patient_info = PatientInfo(
            name="John Doe",
            age=45,
            gender="Male",
            case_number="WC12345",
            injury_date=datetime(2024, 1, 15),
            employer="ABC Company",
            occupation="Construction Worker"
        )
        
        diagnoses = [
            Diagnosis(
                id="diag1",
                icd_code="M54.5",
                description="Lumbar spine strain",
                severity="Moderate"
            ),
            Diagnosis(
                id="diag2",
                icd_code="M75.4",
                description="Right shoulder impingement",
                severity="Mild"
            )
        ]
        
        impairment_ratings = [
            ImpairmentRating(
                id="imp1",
                diagnosis_id="diag1",
                percentage=15,
                ama_table="Table 15-3",
                rationale="Based on AMA Guidelines"
            ),
            ImpairmentRating(
                id="imp2",
                diagnosis_id="diag2",
                percentage=8,
                ama_table="Table 16-3",
                rationale="Based on AMA Guidelines"
            )
        ]
        
        medical_findings = MedicalFindings(
            diagnoses=diagnoses,
            impairment_ratings=impairment_ratings,
            findings=["Limited range of motion in lumbar spine", "Positive impingement sign right shoulder"],
            history="Patient reports onset of back pain following lifting incident at work"
        )
        
        return QMETemplateData(
            patient_info=patient_info,
            medical_findings=medical_findings
        )
    
    @pytest.fixture
    def template_config(self):
        """Create template configuration for testing."""
        return TemplateConfig(
            template_path="test_template.docx",
            output_format="docx",
            apply_professional_formatting=True,
            validate_before_assembly=True,
            validate_after_assembly=True,
            enable_error_recovery=True,
            quality_threshold=70.0,
            max_retry_attempts=3
        )
    
    @pytest.fixture
    def formatting_rules(self):
        """Create formatting rules for testing."""
        return FormattingRules(
            font_family="Times New Roman",
            font_size=12,
            line_spacing=1.15,
            margin_inches=1.0,
            apply_ama_standards=True,
            apply_qme_standards=True
        )
    
    @pytest.fixture
    def quality_requirements(self):
        """Create quality requirements for testing."""
        return QualityRequirements(
            min_completeness_score=80.0,
            min_accuracy_score=85.0,
            min_compliance_score=90.0,
            max_critical_issues=0,
            max_high_issues=2,
            require_all_sections=True,
            require_patient_info=True,
            require_diagnoses=True
        )
    
    @pytest.fixture
    def assembly_engine(self, template_config, formatting_rules, quality_requirements):
        """Create assembly engine instance for testing."""
        return ProfessionalTemplateAssemblyEngine(
            template_config=template_config,
            formatting_rules=formatting_rules,
            quality_requirements=quality_requirements
        )
    
    def test_engine_initialization(self, assembly_engine):
        """Test that the assembly engine initializes correctly."""
        assert assembly_engine is not None
        assert assembly_engine.template_config is not None
        assert assembly_engine.formatting_rules is not None
        assert assembly_engine.quality_requirements is not None
        assert assembly_engine.validator is not None
        assert assembly_engine.formatter is not None
        assert assembly_engine.error_recovery is not None
    
    def test_successful_template_assembly(self, assembly_engine, sample_template_data):
        """Test successful template assembly with complete data."""
        with tempfile.TemporaryDirectory() as temp_dir:
            output_path = os.path.join(temp_dir, "test_output.docx")
            
            result = assembly_engine.assemble_template(
                template_data=sample_template_data,
                output_path=output_path
            )
            
            assert result.status == AssemblyStatus.COMPLETED
            assert result.file_path == output_path
            assert os.path.exists(output_path)
            assert result.file_size_bytes > 0
            assert result.template_data == sample_template_data
            assert result.assembly_time_seconds > 0
    
    def test_template_assembly_with_missing_data(self, assembly_engine):
        """Test template assembly with missing critical data."""
        # Create template data with missing patient name (critical field)
        incomplete_data = QMETemplateData(
            patient_info=PatientInfo(name="", case_number="WC12345"),
            medical_findings=MedicalFindings(diagnoses=[], impairment_ratings=[])
        )
        
        result = assembly_engine.assemble_template(template_data=incomplete_data)
        
        # Should still complete due to error recovery
        assert result.status in [AssemblyStatus.COMPLETED, AssemblyStatus.RECOVERED]
        assert result.file_path is not None
        assert os.path.exists(result.file_path)
    
    def test_pre_assembly_validation(self, assembly_engine, sample_template_data):
        """Test pre-assembly validation functionality."""
        is_valid, issues, quality_score = assembly_engine.validator.validate_pre_assembly(sample_template_data)
        
        assert isinstance(is_valid, bool)
        assert isinstance(issues, list)
        assert isinstance(quality_score, QualityScore)
        assert quality_score.overall_score >= 0
        assert quality_score.overall_score <= 100
    
    def test_post_assembly_validation(self, assembly_engine, sample_template_data):
        """Test post-assembly validation functionality."""
        with tempfile.TemporaryDirectory() as temp_dir:
            output_path = os.path.join(temp_dir, "test_validation.docx")
            
            # First assemble the template
            result = assembly_engine.assemble_template(
                template_data=sample_template_data,
                output_path=output_path
            )
            
            assert result.status == AssemblyStatus.COMPLETED
            
            # Then validate the assembled document
            quality_report = assembly_engine.validate_assembly_quality(result)
            
            assert isinstance(quality_report, QualityReport)
            assert quality_report.overall_score >= 0
            assert quality_report.overall_score <= 100
            assert quality_report.compliance_status in ["compliant", "needs_review", "non_compliant"]
    
    def test_professional_formatting_application(self, assembly_engine, sample_template_data):
        """Test that professional formatting is applied correctly."""
        with tempfile.TemporaryDirectory() as temp_dir:
            output_path = os.path.join(temp_dir, "test_formatting.docx")
            
            result = assembly_engine.assemble_template(
                template_data=sample_template_data,
                output_path=output_path
            )
            
            assert result.status == AssemblyStatus.COMPLETED
            
            # Verify formatting was applied by checking the document
            from docx import Document
            doc = Document(output_path)
            
            # Check that document has content
            assert len(doc.paragraphs) > 0
            
            # Check margins (basic test)
            section = doc.sections[0]
            assert section.top_margin.inches == pytest.approx(1.0, abs=0.1)
    
    def test_error_recovery_mechanisms(self, assembly_engine):
        """Test error recovery with fallback strategies."""
        # Create problematic template data that might cause issues
        problematic_data = QMETemplateData(
            patient_info=PatientInfo(name=None, case_number=None),
            medical_findings=MedicalFindings(diagnoses=None, impairment_ratings=None)
        )
        
        result = assembly_engine.assemble_template(template_data=problematic_data)
        
        # Should recover using fallback strategies
        assert result.status in [AssemblyStatus.COMPLETED, AssemblyStatus.RECOVERED]
        assert result.file_path is not None
        assert os.path.exists(result.file_path)
        
        if result.fallback_used:
            assert result.fallback_used in [
                FallbackStrategy.SIMPLIFIED_TEMPLATE,
                FallbackStrategy.PLACEHOLDER_TEMPLATE,
                FallbackStrategy.MINIMAL_TEMPLATE,
                FallbackStrategy.TEXT_ONLY
            ]
    
    def test_quality_threshold_enforcement(self, template_config, formatting_rules, quality_requirements):
        """Test that quality thresholds are enforced."""
        # Set very high quality requirements
        strict_requirements = QualityRequirements(
            min_completeness_score=95.0,
            min_accuracy_score=95.0,
            min_compliance_score=95.0,
            max_critical_issues=0,
            max_high_issues=0
        )
        
        engine = ProfessionalTemplateAssemblyEngine(
            template_config=template_config,
            formatting_rules=formatting_rules,
            quality_requirements=strict_requirements
        )
        
        # Create data that won't meet strict requirements
        incomplete_data = QMETemplateData(
            patient_info=PatientInfo(name="Test", case_number=""),
            medical_findings=MedicalFindings(diagnoses=[], impairment_ratings=[])
        )
        
        result = engine.assemble_template(template_data=incomplete_data)
        
        # Should still complete due to error recovery being enabled
        assert result.status in [AssemblyStatus.COMPLETED, AssemblyStatus.RECOVERED]
    
    def test_output_path_generation(self, assembly_engine, sample_template_data):
        """Test automatic output path generation."""
        result = assembly_engine.assemble_template(template_data=sample_template_data)
        
        assert result.status == AssemblyStatus.COMPLETED
        assert result.file_path is not None
        assert "John_Doe" in result.file_path or "QME_Report" in result.file_path
        assert result.file_path.endswith(".docx")
        assert os.path.exists(result.file_path)
    
    def test_file_statistics_calculation(self, assembly_engine, sample_template_data):
        """Test that file statistics are calculated correctly."""
        result = assembly_engine.assemble_template(template_data=sample_template_data)
        
        assert result.status == AssemblyStatus.COMPLETED
        assert result.file_size_bytes > 0
        assert result.page_count >= 1
        assert result.word_count > 0
    
    def test_template_content_completeness(self, assembly_engine, sample_template_data):
        """Test that all required content sections are included."""
        result = assembly_engine.assemble_template(template_data=sample_template_data)
        
        assert result.status == AssemblyStatus.COMPLETED
        
        # Read the document and check for required sections
        from docx import Document
        doc = Document(result.file_path)
        doc_text = "\n".join([para.text for para in doc.paragraphs])
        
        # Check for required sections
        required_sections = [
            "QUALIFIED MEDICAL EVALUATOR REPORT",
            "PATIENT IDENTIFICATION",
            "HISTORY OF PRESENT ILLNESS",
            "PHYSICAL EXAMINATION",
            "DIAGNOSIS",
            "IMPAIRMENT RATING"
        ]
        
        for section in required_sections:
            assert section in doc_text, f"Required section '{section}' not found in document"
        
        # Check that patient data is included
        assert sample_template_data.patient_info.name in doc_text
        assert sample_template_data.patient_info.case_number in doc_text
    
    def test_ama_qme_standards_compliance(self, assembly_engine, sample_template_data):
        """Test compliance with AMA and QME standards."""
        result = assembly_engine.assemble_template(template_data=sample_template_data)
        
        assert result.status == AssemblyStatus.COMPLETED
        
        # Validate the quality report for compliance
        quality_report = assembly_engine.validate_assembly_quality(result)
        
        # Should have reasonable compliance scores
        assert quality_report.compliance_score >= 70.0  # Reasonable threshold
        assert quality_report.compliance_status in ["compliant", "needs_review"]
        
        # Should not have critical compliance issues
        assert quality_report.critical_issues == 0
    
    def test_multiple_fallback_strategies(self, template_config, formatting_rules, quality_requirements):
        """Test that multiple fallback strategies work in sequence."""
        # Configure with all fallback strategies
        template_config.fallback_strategies = [
            FallbackStrategy.SIMPLIFIED_TEMPLATE,
            FallbackStrategy.PLACEHOLDER_TEMPLATE,
            FallbackStrategy.MINIMAL_TEMPLATE,
            FallbackStrategy.TEXT_ONLY
        ]
        
        engine = ProfessionalTemplateAssemblyEngine(
            template_config=template_config,
            formatting_rules=formatting_rules,
            quality_requirements=quality_requirements
        )
        
        # Create minimal data that should trigger fallbacks
        minimal_data = QMETemplateData(
            patient_info=PatientInfo(name="Test"),
            medical_findings=MedicalFindings()
        )
        
        result = engine.assemble_template(template_data=minimal_data)
        
        # Should complete with some fallback strategy
        assert result.status in [AssemblyStatus.COMPLETED, AssemblyStatus.RECOVERED]
        assert result.file_path is not None
        assert os.path.exists(result.file_path)
    
    def test_template_validation_rules(self, quality_requirements):
        """Test that validation rules are properly applied."""
        validator = TemplateValidator(quality_requirements)
        
        # Test with complete data
        complete_data = QMETemplateData(
            patient_info=PatientInfo(
                name="John Doe",
                case_number="WC12345",
                injury_date=datetime(2024, 1, 15)
            ),
            medical_findings=MedicalFindings(
                diagnoses=[Diagnosis(id="test1", icd_code="M54.5", description="Test diagnosis")],
                impairment_ratings=[ImpairmentRating(id="imp1", diagnosis_id="test1", percentage=10, ama_table="Table 1", rationale="Test")]
            )
        )
        
        is_valid, issues, quality_score = validator.validate_pre_assembly(complete_data)
        
        assert isinstance(is_valid, bool)
        assert len(issues) >= 0  # May have some minor issues
        assert quality_score.overall_score > 0
        
        # Test with incomplete data
        incomplete_data = QMETemplateData(
            patient_info=PatientInfo(name="", case_number=""),
            medical_findings=MedicalFindings(diagnoses=[], impairment_ratings=[])
        )
        
        is_valid_incomplete, issues_incomplete, quality_score_incomplete = validator.validate_pre_assembly(incomplete_data)
        
        assert len(issues_incomplete) > len(issues)  # Should have more issues
        assert quality_score_incomplete.overall_score < quality_score.overall_score
    
    def test_error_recovery_manager(self):
        """Test the error recovery manager functionality."""
        fallback_strategies = [FallbackStrategy.SIMPLIFIED_TEMPLATE, FallbackStrategy.TEXT_ONLY]
        recovery_manager = ErrorRecoveryManager(fallback_strategies)
        
        # Create test data
        test_data = QMETemplateData(
            patient_info=PatientInfo(name="Test Patient"),
            medical_findings=MedicalFindings()
        )
        
        # Test simplified template recovery
        result = recovery_manager.attempt_recovery(
            template_data=test_data,
            original_error=Exception("Test error"),
            strategy=FallbackStrategy.SIMPLIFIED_TEMPLATE
        )
        
        assert result.status == AssemblyStatus.RECOVERED
        assert result.fallback_used == FallbackStrategy.SIMPLIFIED_TEMPLATE
        assert result.file_path is not None
        assert os.path.exists(result.file_path)
        
        # Clean up
        if result.file_path and os.path.exists(result.file_path):
            os.unlink(result.file_path)
    
    def test_professional_formatter(self, formatting_rules):
        """Test the professional formatter functionality."""
        formatter = ProfessionalFormatter(formatting_rules)
        
        from docx import Document
        doc = Document()
        doc.add_paragraph("Test paragraph")
        
        # Apply formatting
        formatted_doc = formatter.apply_professional_formatting(doc)
        
        assert formatted_doc is not None
        
        # Check that styles were created/updated
        style_names = [style.name for style in formatted_doc.styles]
        assert 'Normal' in style_names
        
        # Check margins
        section = formatted_doc.sections[0]
        expected_margin_inches = formatting_rules.margin_inches
        assert section.top_margin.inches == pytest.approx(expected_margin_inches, abs=0.01)
    
    def test_results_directory_creation(self, assembly_engine, sample_template_data):
        """Test that results directory is created automatically."""
        # Remove results directory if it exists
        results_dir = Path("results/generated_documents")
        if results_dir.exists():
            import shutil
            shutil.rmtree("results")
        
        result = assembly_engine.assemble_template(template_data=sample_template_data)
        
        assert result.status == AssemblyStatus.COMPLETED
        assert results_dir.exists()
        assert result.file_path.startswith(str(results_dir))
    
    def test_assembly_time_tracking(self, assembly_engine, sample_template_data):
        """Test that assembly time is tracked correctly."""
        result = assembly_engine.assemble_template(template_data=sample_template_data)
        
        assert result.status == AssemblyStatus.COMPLETED
        assert result.assembly_time_seconds > 0
        assert result.assembly_time_seconds < 60  # Should complete within reasonable time
    
    def test_quality_report_generation(self, assembly_engine, sample_template_data):
        """Test comprehensive quality report generation."""
        result = assembly_engine.assemble_template(template_data=sample_template_data)
        
        assert result.status == AssemblyStatus.COMPLETED
        
        quality_report = assembly_engine.validate_assembly_quality(result)
        
        # Verify all quality report fields are populated
        assert quality_report.overall_score >= 0
        assert quality_report.completeness_score >= 0
        assert quality_report.accuracy_score >= 0
        assert quality_report.compliance_score >= 0
        assert quality_report.formatting_score >= 0
        assert quality_report.total_issues >= 0
        assert isinstance(quality_report.recommendations, list)
        assert isinstance(quality_report.missing_sections, list)
        assert quality_report.compliance_status in ["compliant", "needs_review", "non_compliant", "error"]
    
    @pytest.mark.parametrize("fallback_strategy", [
        FallbackStrategy.SIMPLIFIED_TEMPLATE,
        FallbackStrategy.PLACEHOLDER_TEMPLATE,
        FallbackStrategy.MINIMAL_TEMPLATE,
        FallbackStrategy.TEXT_ONLY
    ])
    def test_individual_fallback_strategies(self, fallback_strategy):
        """Test each fallback strategy individually."""
        recovery_manager = ErrorRecoveryManager([fallback_strategy])
        
        test_data = QMETemplateData(
            patient_info=PatientInfo(name="Test Patient", case_number="TEST123"),
            medical_findings=MedicalFindings(
                diagnoses=[Diagnosis(id="test1", icd_code="M54.5", description="Test diagnosis")],
                impairment_ratings=[ImpairmentRating(id="imp1", diagnosis_id="test1", percentage=5, ama_table="Table 1", rationale="Test")]
            )
        )
        
        result = recovery_manager.attempt_recovery(
            template_data=test_data,
            original_error=Exception("Test error"),
            strategy=fallback_strategy
        )
        
        assert result.status == AssemblyStatus.RECOVERED
        assert result.fallback_used == fallback_strategy
        assert result.file_path is not None
        assert os.path.exists(result.file_path)
        
        # Verify file has content
        file_size = os.path.getsize(result.file_path)
        assert file_size > 0
        
        # Clean up
        if os.path.exists(result.file_path):
            os.unlink(result.file_path)


class TestTemplateAssemblyIntegration:
    """Integration tests for template assembly with existing system components."""
    
    def test_integration_with_extraction_results(self):
        """Test integration with extraction results from OpenRouter service."""
        # Mock extraction result
        extraction_result = Mock(spec=ExtractionResult)
        extraction_result.confidence_score = 0.85
        extraction_result.extracted_fields = {}
        
        # Create assembly engine
        engine = ProfessionalTemplateAssemblyEngine()
        
        # Create test data
        template_data = QMETemplateData(
            patient_info=PatientInfo(name="Integration Test", case_number="INT123"),
            medical_findings=MedicalFindings()
        )
        
        # Test assembly with extraction result
        result = engine.assemble_template(
            template_data=template_data,
            extraction_result=extraction_result
        )
        
        assert result.status in [AssemblyStatus.COMPLETED, AssemblyStatus.RECOVERED]
        assert result.file_path is not None
    
    def test_template_quality_comparison(self):
        """Test that generated templates meet quality standards of existing successful templates."""
        engine = ProfessionalTemplateAssemblyEngine()
        
        # Create high-quality test data
        high_quality_data = QMETemplateData(
            patient_info=PatientInfo(
                name="Quality Test Patient",
                age=35,
                gender="Female",
                case_number="QT12345",
                injury_date=datetime(2024, 1, 1),
                employer="Quality Corp",
                occupation="Manager"
            ),
            medical_findings=MedicalFindings(
                diagnoses=[
                    Diagnosis(id="diag1", icd_code="M54.2", description="Cervical strain"),
                    Diagnosis(id="diag2", icd_code="G44.1", description="Headaches")
                ],
                impairment_ratings=[
                    ImpairmentRating(id="imp1", diagnosis_id="diag1", percentage=12, ama_table="Table 15-5", rationale="Based on AMA Guidelines")
                ],
                findings=[
                    Finding(id="find1", section_id="exam", finding_type="physical", description="Reduced cervical range of motion", page_reference=1),
                    Finding(id="find2", section_id="exam", finding_type="physical", description="Muscle spasm in upper trapezius", page_reference=1),
                    Finding(id="find3", section_id="exam", finding_type="physical", description="Positive Spurling's test", page_reference=1)
                ]
            )
        )
        
        result = engine.assemble_template(template_data=high_quality_data)
        
        assert result.status == AssemblyStatus.COMPLETED
        
        # Validate quality meets high standards
        quality_report = engine.validate_assembly_quality(result)
        
        # Should meet professional standards
        assert quality_report.overall_score >= 80.0
        assert quality_report.completeness_score >= 85.0
        assert quality_report.compliance_score >= 90.0
        assert quality_report.critical_issues == 0
        assert quality_report.compliance_status in ["compliant", "needs_review"]
    
    def test_template_preservation_and_archiving(self):
        """Test that generated templates are properly preserved in organized folders."""
        engine = ProfessionalTemplateAssemblyEngine()
        
        test_data = QMETemplateData(
            patient_info=PatientInfo(name="Archive Test", case_number="AR123"),
            medical_findings=MedicalFindings()
        )
        
        result = engine.assemble_template(template_data=test_data)
        
        assert result.status in [AssemblyStatus.COMPLETED, AssemblyStatus.RECOVERED]
        
        # Check that file is in results directory structure
        results_path = Path("results/generated_documents")
        assert results_path.exists()
        assert result.file_path.startswith(str(results_path))
        
        # Verify file exists and has content
        assert os.path.exists(result.file_path)
        assert os.path.getsize(result.file_path) > 0


if __name__ == "__main__":
    pytest.main([__file__, "-v"])