"""
Professional Template Assembly Engine - Refactored Version.

This service provides reliable document generation with template validation,
quality checking, professional formatting compliance, and error recovery mechanisms.
Implements requirements 3.1, 3.2, and 3.3 from the QME system refactor specification.
"""

from dataclasses import dataclass, field
from datetime import datetime
from typing import Dict, Any, List, Optional, Tuple, Union
from pathlib import Path
from enum import Enum
import os
import tempfile
import uuid
import json
import yaml

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT

try:
    from src.services.qme_template_generator import QMETemplateData, PatientInfo, MedicalFindings
    from src.services.qme_rules_engine import ValidationIssue, QualityScore, ValidationSeverity
    from src.models.extraction_models import ExtractionResult, ExtractedField
    from src.utils.logging_config import get_logger
except ImportError:
    from services.qme_template_generator import QMETemplateData, PatientInfo, MedicalFindings
    from services.qme_rules_engine import ValidationIssue, QualityScore, ValidationSeverity
    from models.extraction_models import ExtractionResult, ExtractedField
    from utils.logging_config import get_logger

logger = get_logger(__name__)


class AssemblyStatus(Enum):
    """Status of template assembly process."""
    PENDING = "pending"
    VALIDATING = "validating"
    ASSEMBLING = "assembling"
    FORMATTING = "formatting"
    COMPLETED = "completed"
    FAILED = "failed"
    RECOVERED = "recovered"


class FallbackStrategy(Enum):
    """Available fallback strategies for assembly failures."""
    SIMPLIFIED_TEMPLATE = "simplified_template"
    PLACEHOLDER_TEMPLATE = "placeholder_template"
    MINIMAL_TEMPLATE = "minimal_template"
    TEXT_ONLY = "text_only"


@dataclass
class TemplateConfig:
    """Configuration for template assembly."""
    template_path: str
    output_format: str = "docx"
    apply_professional_formatting: bool = True
    validate_before_assembly: bool = True
    validate_after_assembly: bool = True
    enable_error_recovery: bool = True
    fallback_strategies: List[FallbackStrategy] = field(default_factory=lambda: [
        FallbackStrategy.SIMPLIFIED_TEMPLATE,
        FallbackStrategy.PLACEHOLDER_TEMPLATE,
        FallbackStrategy.TEXT_ONLY
    ])
    quality_threshold: float = 70.0
    max_retry_attempts: int = 3


@dataclass
class FormattingRules:
    """Professional formatting rules for QME documents."""
    font_family: str = "Times New Roman"
    font_size: int = 12
    line_spacing: float = 1.15
    margin_inches: float = 1.0
    header_font_size: int = 14
    section_spacing_before: int = 12
    section_spacing_after: int = 6
    paragraph_spacing: int = 6
    apply_ama_standards: bool = True
    apply_qme_standards: bool = True


@dataclass
class QualityRequirements:
    """Quality requirements for template assembly."""
    min_completeness_score: float = 80.0
    min_accuracy_score: float = 85.0
    min_compliance_score: float = 90.0
    max_critical_issues: int = 0
    max_high_issues: int = 2
    require_all_sections: bool = True
    require_patient_info: bool = True
    require_diagnoses: bool = True


@dataclass
class ValidationRule:
    """Individual validation rule."""
    name: str
    description: str
    severity: ValidationSeverity
    check_function: str  # Name of method to call
    auto_fixable: bool = False
    fix_function: Optional[str] = None


@dataclass
class AssemblyResult:
    """Result from template assembly process."""
    status: AssemblyStatus
    file_path: Optional[str] = None
    template_data: Optional[QMETemplateData] = None
    quality_report: Optional['QualityReport'] = None
    validation_issues: List[ValidationIssue] = field(default_factory=list)
    fallback_used: Optional[FallbackStrategy] = None
    assembly_time_seconds: float = 0.0
    file_size_bytes: int = 0
    page_count: int = 0
    word_count: int = 0
    error_message: Optional[str] = None
    recovery_attempts: int = 0


@dataclass
class QualityReport:
    """Comprehensive quality report for assembled template."""
    overall_score: float
    completeness_score: float
    accuracy_score: float
    compliance_score: float
    formatting_score: float
    total_issues: int
    critical_issues: int
    high_issues: int
    medium_issues: int
    low_issues: int
    missing_sections: List[str]
    placeholder_count: int
    compliance_status: str  # "compliant", "needs_review", "non_compliant"
    recommendations: List[str]
    generated_at: datetime = field(default_factory=datetime.now)


class TemplateValidator:
    """Validates templates before and after assembly."""
    
    def __init__(self, quality_requirements: QualityRequirements):
        """Initialize validator with quality requirements."""
        self.quality_requirements = quality_requirements
        self.validation_rules = self._load_validation_rules()
    
    def validate_pre_assembly(self, template_data: QMETemplateData) -> Tuple[bool, List[ValidationIssue], QualityScore]:
        """Validate template data before assembly."""
        try:
            logger.info("Starting pre-assembly validation")
            issues = []
            
            # Validate patient information
            patient_issues = self._validate_patient_info(template_data.patient_info)
            issues.extend(patient_issues)
            
            # Validate medical findings
            medical_issues = self._validate_medical_findings(template_data.medical_findings)
            issues.extend(medical_issues)
            
            # Validate required sections
            section_issues = self._validate_required_sections(template_data)
            issues.extend(section_issues)
            
            # Calculate quality score
            quality_score = self._calculate_quality_score(issues, template_data)
            
            # Determine if validation passes
            is_valid = self._meets_quality_requirements(quality_score, issues)
            
            logger.info(f"Pre-assembly validation complete: {len(issues)} issues, score: {quality_score.overall_score:.1f}")
            return is_valid, issues, quality_score
            
        except Exception as e:
            logger.error(f"Error in pre-assembly validation: {e}")
            return False, [], QualityScore(0, 0, 0, 0, 0)
    
    def validate_post_assembly(self, document_path: str, template_data: QMETemplateData) -> QualityReport:
        """Validate assembled document."""
        try:
            logger.info("Starting post-assembly validation")
            
            # Load document
            doc = Document(document_path)
            
            # Validate document structure
            structure_issues = self._validate_document_structure(doc)
            
            # Validate formatting compliance
            formatting_issues = self._validate_formatting_compliance(doc)
            
            # Check for placeholders
            placeholder_issues = self._check_remaining_placeholders(doc)
            
            # Validate content completeness
            completeness_issues = self._validate_content_completeness(doc, template_data)
            
            # Combine all issues
            all_issues = structure_issues + formatting_issues + placeholder_issues + completeness_issues
            
            # Calculate scores
            quality_scores = self._calculate_post_assembly_scores(doc, all_issues)
            
            # Generate quality report
            quality_report = self._generate_quality_report(doc, all_issues, quality_scores)
            
            logger.info(f"Post-assembly validation complete: {quality_report.total_issues} issues")
            return quality_report
            
        except Exception as e:
            logger.error(f"Error in post-assembly validation: {e}")
            return self._create_error_quality_report(str(e))
    
    def _load_validation_rules(self) -> List[ValidationRule]:
        """Load validation rules from configuration."""
        try:
            # Load from config file if available
            config_path = Path("config/prompts/validation/compliance_rules.yaml")
            if config_path.exists():
                with open(config_path, 'r') as f:
                    config = yaml.safe_load(f)
                    rules = []
                    for rule_data in config.get('validation_rules', []):
                        if isinstance(rule_data, dict):
                            rules.append(ValidationRule(**rule_data))
                    return rules
            
            # Default validation rules
            return [
                ValidationRule(
                    name="patient_name_required",
                    description="Patient name must be present",
                    severity=ValidationSeverity.CRITICAL,
                    check_function="_check_patient_name"
                ),
                ValidationRule(
                    name="case_number_required",
                    description="Case number must be present",
                    severity=ValidationSeverity.HIGH,
                    check_function="_check_case_number"
                ),
                ValidationRule(
                    name="diagnosis_required",
                    description="At least one diagnosis must be present",
                    severity=ValidationSeverity.CRITICAL,
                    check_function="_check_diagnoses"
                ),
                ValidationRule(
                    name="impairment_rating_required",
                    description="Impairment rating must be present",
                    severity=ValidationSeverity.HIGH,
                    check_function="_check_impairment_rating"
                )
            ]
            
        except Exception as e:
            logger.error(f"Error loading validation rules: {e}")
            return []
    
    def _validate_patient_info(self, patient_info: PatientInfo) -> List[ValidationIssue]:
        """Validate patient information completeness."""
        issues = []
        
        if not patient_info.name:
            issues.append(ValidationIssue(
                section="patient_info",
                severity=ValidationSeverity.CRITICAL,
                title="Missing Patient Name",
                description="Patient name is required for QME report",
                suggestions=["Ensure patient name is extracted from source documents"]
            ))
        
        if not patient_info.case_number:
            issues.append(ValidationIssue(
                section="patient_info",
                severity=ValidationSeverity.HIGH,
                title="Missing Case Number",
                description="Case number is required for proper identification",
                suggestions=["Extract case number from claim documents"]
            ))
        
        if not patient_info.injury_date:
            issues.append(ValidationIssue(
                section="patient_info",
                severity=ValidationSeverity.HIGH,
                title="Missing Injury Date",
                description="Date of injury is required for QME evaluation",
                suggestions=["Extract injury date from medical records"]
            ))
        
        return issues
    
    def _validate_medical_findings(self, medical_findings: MedicalFindings) -> List[ValidationIssue]:
        """Validate medical findings completeness."""
        issues = []
        
        if not medical_findings.diagnoses:
            issues.append(ValidationIssue(
                section="medical_findings",
                severity=ValidationSeverity.CRITICAL,
                title="Missing Diagnoses",
                description="At least one diagnosis is required",
                suggestions=["Extract diagnoses from medical examination records"]
            ))
        
        if not medical_findings.impairment_ratings:
            issues.append(ValidationIssue(
                section="medical_findings",
                severity=ValidationSeverity.HIGH,
                title="Missing Impairment Rating",
                description="Impairment rating is required for QME report",
                suggestions=["Calculate impairment rating based on AMA Guidelines"]
            ))
        
        return issues
    
    def _validate_required_sections(self, template_data: QMETemplateData) -> List[ValidationIssue]:
        """Validate that all required sections have content."""
        issues = []
        
        # Check for examination findings
        if not template_data.medical_findings.findings:
            issues.append(ValidationIssue(
                section="examination",
                severity=ValidationSeverity.MEDIUM,
                title="Missing Examination Findings",
                description="Physical examination findings are recommended",
                suggestions=["Include detailed examination findings"]
            ))
        
        return issues
    
    def _calculate_quality_score(self, issues: List[ValidationIssue], template_data: QMETemplateData) -> QualityScore:
        """Calculate quality score based on validation issues."""
        # Count issues by severity
        critical_count = len([i for i in issues if i.severity == ValidationSeverity.CRITICAL])
        high_count = len([i for i in issues if i.severity == ValidationSeverity.HIGH])
        medium_count = len([i for i in issues if i.severity == ValidationSeverity.MEDIUM])
        low_count = len([i for i in issues if i.severity == ValidationSeverity.LOW])
        
        # Calculate component scores
        completeness_score = max(0, 100 - (critical_count * 30) - (high_count * 15) - (medium_count * 5))
        accuracy_score = max(0, 100 - (critical_count * 25) - (high_count * 10))
        compliance_score = max(0, 100 - (critical_count * 40) - (high_count * 20) - (medium_count * 5))
        
        overall_score = (completeness_score + accuracy_score + compliance_score) / 3
        
        return QualityScore(
            overall_score=overall_score,
            completeness_score=completeness_score,
            accuracy_score=accuracy_score,
            compliance_score=compliance_score,
            evidence_confidence_score=0.0,  # Not calculated in pre-assembly
            total_issues=len(issues),
            critical_issues=critical_count,
            high_issues=high_count,
            medium_issues=medium_count,
            low_issues=low_count
        )
    
    def _meets_quality_requirements(self, quality_score: QualityScore, issues: List[ValidationIssue]) -> bool:
        """Check if quality score meets requirements."""
        critical_issues = len([i for i in issues if i.severity == ValidationSeverity.CRITICAL])
        high_issues = len([i for i in issues if i.severity == ValidationSeverity.HIGH])
        
        return (
            quality_score.completeness_score >= self.quality_requirements.min_completeness_score and
            quality_score.accuracy_score >= self.quality_requirements.min_accuracy_score and
            quality_score.compliance_score >= self.quality_requirements.min_compliance_score and
            critical_issues <= self.quality_requirements.max_critical_issues and
            high_issues <= self.quality_requirements.max_high_issues
        )
    
    def _validate_document_structure(self, doc: Document) -> List[ValidationIssue]:
        """Validate document structure against QME standards."""
        issues = []
        
        # Get all paragraph text
        doc_text = "\n".join([para.text for para in doc.paragraphs])
        
        # Check for required sections
        required_sections = [
            "PATIENT IDENTIFICATION",
            "HISTORY OF PRESENT ILLNESS", 
            "PHYSICAL EXAMINATION",
            "DIAGNOSIS",
            "IMPAIRMENT RATING"
        ]
        
        for section in required_sections:
            if section not in doc_text.upper():
                issues.append(ValidationIssue(
                    section="document_structure",
                    severity=ValidationSeverity.HIGH,
                    title=f"Missing Required Section: {section}",
                    description=f"Required section '{section}' not found in document",
                    suggestions=[f"Add {section} section to document"]
                ))
        
        return issues
    
    def _validate_formatting_compliance(self, doc: Document) -> List[ValidationIssue]:
        """Validate formatting compliance with professional standards."""
        issues = []
        
        try:
            # Check margins
            section = doc.sections[0]
            expected_margin = Inches(1.0)
            
            if abs(section.top_margin - expected_margin) > Inches(0.1):
                issues.append(ValidationIssue(
                    section="formatting",
                    severity=ValidationSeverity.LOW,
                    title="Incorrect Margin",
                    description=f"Top margin is {section.top_margin.inches:.1f} inches, expected 1.0 inches",
                    suggestions=["Adjust margins to 1.0 inches"],
                    auto_fixable=True
                ))
            
            # Check font consistency (basic check)
            inconsistent_fonts = 0
            for para in doc.paragraphs:
                for run in para.runs:
                    if run.font.name and run.font.name != "Times New Roman":
                        inconsistent_fonts += 1
            
            if inconsistent_fonts > 0:
                issues.append(ValidationIssue(
                    section="formatting",
                    severity=ValidationSeverity.MEDIUM,
                    title="Inconsistent Font Usage",
                    description=f"Found {inconsistent_fonts} text runs with non-standard fonts",
                    suggestions=["Apply consistent Times New Roman font throughout document"],
                    auto_fixable=True
                ))
                
        except Exception as e:
            logger.error(f"Error validating formatting: {e}")
        
        return issues
    
    def _check_remaining_placeholders(self, doc: Document) -> List[ValidationIssue]:
        """Check for remaining placeholder text."""
        issues = []
        
        doc_text = "\n".join([para.text for para in doc.paragraphs])
        
        # Common placeholder patterns
        import re
        placeholder_patterns = [
            r'\[MISSING[^\]]*\]',
            r'\[ERROR[^\]]*\]',
            r'\[TODO[^\]]*\]',
            r'\+[A-Za-z\s]+\+',  # Template placeholders like +Patient Name+
            r'_{3,}',  # Multiple underscores
        ]
        
        for pattern in placeholder_patterns:
            matches = re.findall(pattern, doc_text)
            for match in matches:
                issues.append(ValidationIssue(
                    section="content_completeness",
                    severity=ValidationSeverity.MEDIUM,
                    title="Placeholder Text Found",
                    description=f"Placeholder text found: {match}",
                    suggestions=["Replace placeholder with actual content"]
                ))
        
        return issues
    
    def _validate_content_completeness(self, doc: Document, template_data: QMETemplateData) -> List[ValidationIssue]:
        """Validate content completeness against template data."""
        issues = []
        
        # This is a simplified check - in full implementation would be more comprehensive
        doc_text = "\n".join([para.text for para in doc.paragraphs])
        
        # Check if patient name appears in document
        if template_data.patient_info.name and template_data.patient_info.name not in doc_text:
            issues.append(ValidationIssue(
                section="content_completeness",
                severity=ValidationSeverity.HIGH,
                title="Patient Name Not Found in Document",
                description="Patient name from template data not found in assembled document",
                suggestions=["Ensure patient name is properly inserted into template"]
            ))
        
        return issues
    
    def _calculate_post_assembly_scores(self, doc: Document, issues: List[ValidationIssue]) -> Dict[str, float]:
        """Calculate quality scores for assembled document."""
        critical_count = len([i for i in issues if i.severity == ValidationSeverity.CRITICAL])
        high_count = len([i for i in issues if i.severity == ValidationSeverity.HIGH])
        medium_count = len([i for i in issues if i.severity == ValidationSeverity.MEDIUM])
        low_count = len([i for i in issues if i.severity == ValidationSeverity.LOW])
        
        # Calculate component scores
        completeness_score = max(0, 100 - (critical_count * 25) - (high_count * 10) - (medium_count * 5))
        accuracy_score = max(0, 100 - (critical_count * 30) - (high_count * 15))
        compliance_score = max(0, 100 - (critical_count * 40) - (high_count * 20) - (medium_count * 5))
        formatting_score = max(0, 100 - (low_count * 2) - (medium_count * 5))
        
        overall_score = (completeness_score + accuracy_score + compliance_score + formatting_score) / 4
        
        return {
            'overall': overall_score,
            'completeness': completeness_score,
            'accuracy': accuracy_score,
            'compliance': compliance_score,
            'formatting': formatting_score
        }
    
    def _generate_quality_report(self, doc: Document, issues: List[ValidationIssue], scores: Dict[str, float]) -> QualityReport:
        """Generate comprehensive quality report."""
        # Count issues by severity
        critical_count = len([i for i in issues if i.severity == ValidationSeverity.CRITICAL])
        high_count = len([i for i in issues if i.severity == ValidationSeverity.HIGH])
        medium_count = len([i for i in issues if i.severity == ValidationSeverity.MEDIUM])
        low_count = len([i for i in issues if i.severity == ValidationSeverity.LOW])
        
        # Identify missing sections
        missing_sections = []
        for issue in issues:
            if "Missing Required Section" in issue.title:
                missing_sections.append(issue.title.split(": ")[1])
        
        # Count placeholders
        placeholder_count = len([i for i in issues if "Placeholder" in i.title])
        
        # Determine compliance status
        if critical_count > 0:
            compliance_status = "non_compliant"
        elif high_count > 2 or scores['overall'] < 70:
            compliance_status = "needs_review"
        else:
            compliance_status = "compliant"
        
        # Generate recommendations
        recommendations = self._generate_recommendations(issues, scores)
        
        return QualityReport(
            overall_score=scores['overall'],
            completeness_score=scores['completeness'],
            accuracy_score=scores['accuracy'],
            compliance_score=scores['compliance'],
            formatting_score=scores['formatting'],
            total_issues=len(issues),
            critical_issues=critical_count,
            high_issues=high_count,
            medium_issues=medium_count,
            low_issues=low_count,
            missing_sections=missing_sections,
            placeholder_count=placeholder_count,
            compliance_status=compliance_status,
            recommendations=recommendations
        )
    
    def _generate_recommendations(self, issues: List[ValidationIssue], scores: Dict[str, float]) -> List[str]:
        """Generate improvement recommendations."""
        recommendations = []
        
        if scores['completeness'] < 80:
            recommendations.append("Improve data extraction to ensure all required fields are populated")
        
        if scores['accuracy'] < 85:
            recommendations.append("Review and validate extracted information for accuracy")
        
        if scores['compliance'] < 90:
            recommendations.append("Ensure all QME regulatory requirements are met")
        
        if scores['formatting'] < 95:
            recommendations.append("Apply consistent professional formatting throughout document")
        
        # Add specific recommendations based on issues
        critical_issues = [i for i in issues if i.severity == ValidationSeverity.CRITICAL]
        if critical_issues:
            recommendations.append("Address all critical issues before finalizing document")
        
        return recommendations
    
    def _create_error_quality_report(self, error_message: str) -> QualityReport:
        """Create quality report for validation errors."""
        return QualityReport(
            overall_score=0.0,
            completeness_score=0.0,
            accuracy_score=0.0,
            compliance_score=0.0,
            formatting_score=0.0,
            total_issues=1,
            critical_issues=1,
            high_issues=0,
            medium_issues=0,
            low_issues=0,
            missing_sections=[],
            placeholder_count=0,
            compliance_status="error",
            recommendations=[f"Fix validation error: {error_message}"]
        )


class ProfessionalFormatter:
    """Handles professional formatting according to AMA and QME standards."""
    
    def __init__(self, formatting_rules: FormattingRules):
        """Initialize formatter with formatting rules."""
        self.formatting_rules = formatting_rules
    
    def apply_professional_formatting(self, doc: Document) -> Document:
        """Apply professional formatting to document."""
        try:
            logger.info("Applying professional formatting")
            
            # Apply document-level formatting
            self._apply_document_formatting(doc)
            
            # Create and apply styles
            self._create_professional_styles(doc)
            
            # Apply paragraph formatting
            self._apply_paragraph_formatting(doc)
            
            # Apply header and footer
            self._apply_header_footer(doc)
            
            logger.info("Professional formatting applied successfully")
            return doc
            
        except Exception as e:
            logger.error(f"Error applying professional formatting: {e}")
            raise
    
    def _apply_document_formatting(self, doc: Document) -> None:
        """Apply document-level formatting settings."""
        # Set margins
        for section in doc.sections:
            margin = Inches(self.formatting_rules.margin_inches)
            section.top_margin = margin
            section.bottom_margin = margin
            section.left_margin = margin
            section.right_margin = margin
    
    def _create_professional_styles(self, doc: Document) -> None:
        """Create professional styles for QME documents."""
        styles = doc.styles
        
        # Update Normal style
        normal_style = styles['Normal']
        normal_font = normal_style.font
        normal_font.name = self.formatting_rules.font_family
        normal_font.size = Pt(self.formatting_rules.font_size)
        
        normal_paragraph = normal_style.paragraph_format
        normal_paragraph.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        normal_paragraph.line_spacing = self.formatting_rules.line_spacing
        normal_paragraph.space_after = Pt(self.formatting_rules.paragraph_spacing)
        
        # Create QME Header style
        try:
            header_style = styles.add_style('QME Header', WD_STYLE_TYPE.PARAGRAPH)
            header_font = header_style.font
            header_font.name = self.formatting_rules.font_family
            header_font.size = Pt(self.formatting_rules.header_font_size)
            header_font.bold = True
            
            header_paragraph = header_style.paragraph_format
            header_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            header_paragraph.space_before = Pt(self.formatting_rules.section_spacing_before)
            header_paragraph.space_after = Pt(self.formatting_rules.section_spacing_after)
        except:
            pass  # Style might already exist
        
        # Create QME Section style
        try:
            section_style = styles.add_style('QME Section', WD_STYLE_TYPE.PARAGRAPH)
            section_font = section_style.font
            section_font.name = self.formatting_rules.font_family
            section_font.size = Pt(self.formatting_rules.font_size)
            section_font.bold = True
            
            section_paragraph = section_style.paragraph_format
            section_paragraph.space_before = Pt(self.formatting_rules.section_spacing_before)
            section_paragraph.space_after = Pt(self.formatting_rules.section_spacing_after)
        except:
            pass  # Style might already exist
    
    def _apply_paragraph_formatting(self, doc: Document) -> None:
        """Apply consistent paragraph formatting."""
        for paragraph in doc.paragraphs:
            # Apply font formatting to all runs
            for run in paragraph.runs:
                if not run.font.name:
                    run.font.name = self.formatting_rules.font_family
                if not run.font.size:
                    run.font.size = Pt(self.formatting_rules.font_size)
    
    def _apply_header_footer(self, doc: Document) -> None:
        """Apply professional header and footer."""
        try:
            section = doc.sections[0]
            
            # Header
            header = section.header
            if header.paragraphs:
                header_para = header.paragraphs[0]
                header_para.text = "QUALIFIED MEDICAL EVALUATOR REPORT"
                header_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                
                if header_para.runs:
                    header_run = header_para.runs[0]
                    header_run.font.name = self.formatting_rules.font_family
                    header_run.font.size = Pt(10)
                    header_run.font.bold = True
            
            # Footer
            footer = section.footer
            if footer.paragraphs:
                footer_para = footer.paragraphs[0]
                footer_para.text = "CONFIDENTIAL MEDICAL REPORT"
                footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                
                if footer_para.runs:
                    footer_run = footer_para.runs[0]
                    footer_run.font.name = self.formatting_rules.font_family
                    footer_run.font.size = Pt(9)
                    footer_run.font.italic = True
                    
        except Exception as e:
            logger.error(f"Error applying header/footer: {e}")


class ErrorRecoveryManager:
    """Manages error recovery and fallback strategies."""
    
    def __init__(self, fallback_strategies: List[FallbackStrategy]):
        """Initialize error recovery manager."""
        self.fallback_strategies = fallback_strategies
        self.recovery_attempts = 0
        self.max_attempts = 3
    
    def attempt_recovery(self, 
                        template_data: QMETemplateData, 
                        original_error: Exception,
                        strategy: FallbackStrategy) -> AssemblyResult:
        """Attempt recovery using specified fallback strategy."""
        try:
            self.recovery_attempts += 1
            logger.info(f"Attempting recovery with strategy: {strategy.value} (attempt {self.recovery_attempts})")
            
            if strategy == FallbackStrategy.SIMPLIFIED_TEMPLATE:
                return self._create_simplified_template(template_data)
            elif strategy == FallbackStrategy.PLACEHOLDER_TEMPLATE:
                return self._create_placeholder_template(template_data)
            elif strategy == FallbackStrategy.MINIMAL_TEMPLATE:
                return self._create_minimal_template(template_data)
            elif strategy == FallbackStrategy.TEXT_ONLY:
                return self._create_text_only_template(template_data)
            else:
                raise ValueError(f"Unknown fallback strategy: {strategy}")
                
        except Exception as e:
            logger.error(f"Recovery attempt failed with strategy {strategy.value}: {e}")
            return AssemblyResult(
                status=AssemblyStatus.FAILED,
                error_message=f"Recovery failed: {str(e)}",
                recovery_attempts=self.recovery_attempts
            )
    
    def _create_simplified_template(self, template_data: QMETemplateData) -> AssemblyResult:
        """Create simplified template with basic formatting."""
        try:
            # Create temporary file
            temp_file = tempfile.NamedTemporaryFile(delete=False, suffix='.docx')
            temp_path = temp_file.name
            temp_file.close()
            
            # Create simple document
            doc = Document()
            
            # Add title
            title = doc.add_paragraph("QUALIFIED MEDICAL EVALUATOR REPORT")
            title.alignment = WD_ALIGN_PARAGRAPH.CENTER
            title.runs[0].font.bold = True
            title.runs[0].font.size = Pt(16)
            
            # Add patient info
            doc.add_paragraph()
            patient_section = doc.add_paragraph("PATIENT INFORMATION")
            patient_section.runs[0].font.bold = True
            
            doc.add_paragraph(f"Name: {template_data.patient_info.name or '[MISSING]'}")
            doc.add_paragraph(f"Case Number: {template_data.patient_info.case_number or '[MISSING]'}")
            doc.add_paragraph(f"Date of Injury: {template_data.patient_info.injury_date or '[MISSING]'}")
            
            # Add diagnoses
            doc.add_paragraph()
            diagnosis_section = doc.add_paragraph("DIAGNOSES")
            diagnosis_section.runs[0].font.bold = True
            
            if template_data.medical_findings.diagnoses:
                for i, diagnosis in enumerate(template_data.medical_findings.diagnoses, 1):
                    doc.add_paragraph(f"{i}. {diagnosis.description}")
            else:
                doc.add_paragraph("[NO DIAGNOSES AVAILABLE]")
            
            # Add impairment ratings
            doc.add_paragraph()
            impairment_section = doc.add_paragraph("IMPAIRMENT RATING")
            impairment_section.runs[0].font.bold = True
            
            if template_data.medical_findings.impairment_ratings:
                for rating in template_data.medical_findings.impairment_ratings:
                    doc.add_paragraph(f"- {rating.percentage}% impairment")
            else:
                doc.add_paragraph("[NO IMPAIRMENT RATING AVAILABLE]")
            
            # Save document
            doc.save(temp_path)
            
            # Get file stats
            file_size = os.path.getsize(temp_path)
            
            return AssemblyResult(
                status=AssemblyStatus.RECOVERED,
                file_path=temp_path,
                template_data=template_data,
                fallback_used=FallbackStrategy.SIMPLIFIED_TEMPLATE,
                file_size_bytes=file_size,
                page_count=1,
                word_count=len(doc.paragraphs) * 5,  # Rough estimate
                recovery_attempts=self.recovery_attempts
            )
            
        except Exception as e:
            logger.error(f"Failed to create simplified template: {e}")
            raise
    
    def _create_placeholder_template(self, template_data: QMETemplateData) -> AssemblyResult:
        """Create template with placeholders for missing data."""
        try:
            # Similar to simplified but with more placeholders
            temp_file = tempfile.NamedTemporaryFile(delete=False, suffix='.docx')
            temp_path = temp_file.name
            temp_file.close()
            
            doc = Document()
            
            # Add comprehensive placeholder template
            doc.add_paragraph("QUALIFIED MEDICAL EVALUATOR REPORT").alignment = WD_ALIGN_PARAGRAPH.CENTER
            doc.add_paragraph()
            
            # Patient information with placeholders
            doc.add_paragraph("PATIENT IDENTIFICATION").runs[0].font.bold = True
            doc.add_paragraph(f"Patient Name: {template_data.patient_info.name or '[PATIENT NAME REQUIRED]'}")
            doc.add_paragraph(f"Date of Birth: {template_data.patient_info.date_of_birth or '[DATE OF BIRTH REQUIRED]'}")
            doc.add_paragraph(f"Case Number: {template_data.patient_info.case_number or '[CASE NUMBER REQUIRED]'}")
            doc.add_paragraph(f"Date of Injury: {template_data.patient_info.injury_date or '[INJURY DATE REQUIRED]'}")
            doc.add_paragraph(f"Employer: {template_data.patient_info.employer or '[EMPLOYER REQUIRED]'}")
            
            doc.add_paragraph()
            doc.add_paragraph("HISTORY OF PRESENT ILLNESS").runs[0].font.bold = True
            doc.add_paragraph("[HISTORY SECTION REQUIRES MANUAL COMPLETION]")
            
            doc.add_paragraph()
            doc.add_paragraph("PHYSICAL EXAMINATION").runs[0].font.bold = True
            doc.add_paragraph("[EXAMINATION FINDINGS REQUIRE MANUAL COMPLETION]")
            
            doc.add_paragraph()
            doc.add_paragraph("DIAGNOSIS").runs[0].font.bold = True
            if template_data.medical_findings.diagnoses:
                for i, diagnosis in enumerate(template_data.medical_findings.diagnoses, 1):
                    doc.add_paragraph(f"{i}. {diagnosis.description}")
            else:
                doc.add_paragraph("[DIAGNOSES REQUIRE MANUAL COMPLETION]")
            
            doc.add_paragraph()
            doc.add_paragraph("IMPAIRMENT RATING").runs[0].font.bold = True
            if template_data.medical_findings.impairment_ratings:
                for rating in template_data.medical_findings.impairment_ratings:
                    doc.add_paragraph(f"- {rating.percentage}% impairment")
            else:
                doc.add_paragraph("[IMPAIRMENT RATING REQUIRES MANUAL COMPLETION]")
            
            doc.save(temp_path)
            file_size = os.path.getsize(temp_path)
            
            return AssemblyResult(
                status=AssemblyStatus.RECOVERED,
                file_path=temp_path,
                template_data=template_data,
                fallback_used=FallbackStrategy.PLACEHOLDER_TEMPLATE,
                file_size_bytes=file_size,
                page_count=2,
                word_count=len(doc.paragraphs) * 8,
                recovery_attempts=self.recovery_attempts
            )
            
        except Exception as e:
            logger.error(f"Failed to create placeholder template: {e}")
            raise
    
    def _create_minimal_template(self, template_data: QMETemplateData) -> AssemblyResult:
        """Create minimal template with only essential information."""
        try:
            temp_file = tempfile.NamedTemporaryFile(delete=False, suffix='.docx')
            temp_path = temp_file.name
            temp_file.close()
            
            doc = Document()
            
            # Minimal content
            doc.add_paragraph("QME REPORT - MINIMAL VERSION")
            doc.add_paragraph(f"Patient: {template_data.patient_info.name or 'Unknown'}")
            doc.add_paragraph(f"Case: {template_data.patient_info.case_number or 'Unknown'}")
            
            if template_data.medical_findings.diagnoses:
                doc.add_paragraph(f"Primary Diagnosis: {template_data.medical_findings.diagnoses[0].description}")
            
            if template_data.medical_findings.impairment_ratings:
                doc.add_paragraph(f"Impairment: {template_data.medical_findings.impairment_ratings[0].percentage}%")
            
            doc.save(temp_path)
            file_size = os.path.getsize(temp_path)
            
            return AssemblyResult(
                status=AssemblyStatus.RECOVERED,
                file_path=temp_path,
                template_data=template_data,
                fallback_used=FallbackStrategy.MINIMAL_TEMPLATE,
                file_size_bytes=file_size,
                page_count=1,
                word_count=30,
                recovery_attempts=self.recovery_attempts
            )
            
        except Exception as e:
            logger.error(f"Failed to create minimal template: {e}")
            raise
    
    def _create_text_only_template(self, template_data: QMETemplateData) -> AssemblyResult:
        """Create text-only template as last resort."""
        try:
            temp_file = tempfile.NamedTemporaryFile(delete=False, suffix='.txt', mode='w')
            temp_path = temp_file.name
            
            # Write text content
            temp_file.write("QME REPORT - TEXT VERSION\n")
            temp_file.write("=" * 40 + "\n\n")
            temp_file.write(f"Patient: {template_data.patient_info.name or 'Unknown'}\n")
            temp_file.write(f"Case: {template_data.patient_info.case_number or 'Unknown'}\n")
            temp_file.write(f"Injury Date: {template_data.patient_info.injury_date or 'Unknown'}\n\n")
            
            temp_file.write("DIAGNOSES:\n")
            if template_data.medical_findings.diagnoses:
                for i, diagnosis in enumerate(template_data.medical_findings.diagnoses, 1):
                    temp_file.write(f"{i}. {diagnosis.description}\n")
            else:
                temp_file.write("No diagnoses available\n")
            
            temp_file.write("\nIMPAIRMENT RATINGS:\n")
            if template_data.medical_findings.impairment_ratings:
                for rating in template_data.medical_findings.impairment_ratings:
                    temp_file.write(f"- {rating.percentage}% impairment\n")
            else:
                temp_file.write("No impairment ratings available\n")
            
            temp_file.close()
            
            file_size = os.path.getsize(temp_path)
            
            return AssemblyResult(
                status=AssemblyStatus.RECOVERED,
                file_path=temp_path,
                template_data=template_data,
                fallback_used=FallbackStrategy.TEXT_ONLY,
                file_size_bytes=file_size,
                page_count=1,
                word_count=50,
                recovery_attempts=self.recovery_attempts
            )
            
        except Exception as e:
            logger.error(f"Failed to create text-only template: {e}")
            raise


class ProfessionalTemplateAssemblyEngine:
    """
    Professional Template Assembly Engine with reliable document generation,
    template validation, quality checking, and error recovery mechanisms.
    
    Implements requirements 3.1, 3.2, and 3.3 from QME system refactor specification.
    """
    
    def __init__(self, 
                 template_config: Optional[TemplateConfig] = None,
                 formatting_rules: Optional[FormattingRules] = None,
                 quality_requirements: Optional[QualityRequirements] = None):
        """
        Initialize the Professional Template Assembly Engine.
        
        Args:
            template_config: Configuration for template assembly
            formatting_rules: Professional formatting rules
            quality_requirements: Quality requirements for validation
        """
        self.template_config = template_config or TemplateConfig(template_path="")
        self.formatting_rules = formatting_rules or FormattingRules()
        self.quality_requirements = quality_requirements or QualityRequirements()
        
        # Initialize components
        self.validator = TemplateValidator(self.quality_requirements)
        self.formatter = ProfessionalFormatter(self.formatting_rules)
        self.error_recovery = ErrorRecoveryManager(self.template_config.fallback_strategies)
        
        logger.info("Professional Template Assembly Engine initialized")
    
    def assemble_template(self, 
                         template_data: QMETemplateData,
                         extraction_result: Optional[ExtractionResult] = None,
                         output_path: Optional[str] = None) -> AssemblyResult:
        """
        Assemble professional template with validation and error recovery.
        
        Args:
            template_data: QME template data to assemble
            extraction_result: Optional extraction result for enhanced validation
            output_path: Optional output file path
            
        Returns:
            AssemblyResult with assembly status and details
        """
        start_time = datetime.now()
        
        try:
            logger.info("Starting professional template assembly")
            
            # Pre-assembly validation
            if self.template_config.validate_before_assembly:
                is_valid, issues, quality_score = self.validator.validate_pre_assembly(template_data)
                
                if not is_valid and not self.template_config.enable_error_recovery:
                    return AssemblyResult(
                        status=AssemblyStatus.FAILED,
                        template_data=template_data,
                        validation_issues=issues,
                        error_message="Pre-assembly validation failed",
                        assembly_time_seconds=(datetime.now() - start_time).total_seconds()
                    )
            
            # Generate output path if not provided
            if not output_path:
                output_path = self._generate_output_path(template_data)
            
            # Attempt main assembly
            try:
                result = self._perform_main_assembly(template_data, output_path)
                
                # Post-assembly validation
                if self.template_config.validate_after_assembly and result.file_path:
                    quality_report = self.validator.validate_post_assembly(result.file_path, template_data)
                    result.quality_report = quality_report
                    
                    # Check if quality meets requirements
                    if quality_report.compliance_status == "non_compliant" and self.template_config.enable_error_recovery:
                        logger.warning("Assembly quality below requirements, attempting recovery")
                        return self._attempt_error_recovery(template_data, Exception("Quality below requirements"))
                
                result.assembly_time_seconds = (datetime.now() - start_time).total_seconds()
                logger.info(f"Template assembly completed successfully in {result.assembly_time_seconds:.2f} seconds")
                return result
                
            except Exception as e:
                logger.error(f"Main assembly failed: {e}")
                
                if self.template_config.enable_error_recovery:
                    return self._attempt_error_recovery(template_data, e)
                else:
                    raise
                    
        except Exception as e:
            logger.error(f"Template assembly failed: {e}")
            return AssemblyResult(
                status=AssemblyStatus.FAILED,
                template_data=template_data,
                error_message=str(e),
                assembly_time_seconds=(datetime.now() - start_time).total_seconds()
            )
    
    def _perform_main_assembly(self, template_data: QMETemplateData, output_path: str) -> AssemblyResult:
        """Perform the main template assembly process."""
        try:
            # Create document
            doc = Document()
            
            # Apply professional formatting
            if self.template_config.apply_professional_formatting:
                doc = self.formatter.apply_professional_formatting(doc)
            
            # Assemble content
            self._assemble_document_content(doc, template_data)
            
            # Save document
            doc.save(output_path)
            
            # Get file statistics
            file_stats = self._get_file_statistics(output_path)
            
            return AssemblyResult(
                status=AssemblyStatus.COMPLETED,
                file_path=output_path,
                template_data=template_data,
                file_size_bytes=file_stats['size'],
                page_count=file_stats['pages'],
                word_count=file_stats['words']
            )
            
        except Exception as e:
            logger.error(f"Main assembly failed: {e}")
            raise
    
    def _assemble_document_content(self, doc: Document, template_data: QMETemplateData) -> None:
        """Assemble the main document content."""
        # Document title
        title = doc.add_paragraph("QUALIFIED MEDICAL EVALUATOR REPORT")
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if title.runs:
            title.runs[0].font.bold = True
            title.runs[0].font.size = Pt(16)
        
        # Patient identification section
        doc.add_paragraph()
        patient_header = doc.add_paragraph("PATIENT IDENTIFICATION")
        if patient_header.runs:
            patient_header.runs[0].font.bold = True
            patient_header.runs[0].font.size = Pt(14)
        
        doc.add_paragraph(f"Patient Name: {template_data.patient_info.name or '[MISSING PATIENT NAME]'}")
        doc.add_paragraph(f"Date of Birth: {getattr(template_data.patient_info, 'date_of_birth', None) or '[MISSING DATE OF BIRTH]'}")
        doc.add_paragraph(f"Age: {template_data.patient_info.age or '[MISSING AGE]'}")
        doc.add_paragraph(f"Gender: {template_data.patient_info.gender or '[MISSING GENDER]'}")
        doc.add_paragraph(f"Case Number: {template_data.patient_info.case_number or '[MISSING CASE NUMBER]'}")
        doc.add_paragraph(f"Date of Injury: {template_data.patient_info.injury_date or '[MISSING INJURY DATE]'}")
        doc.add_paragraph(f"Employer: {template_data.patient_info.employer or '[MISSING EMPLOYER]'}")
        doc.add_paragraph(f"Occupation: {template_data.patient_info.occupation or '[MISSING OCCUPATION]'}")
        
        # History section
        doc.add_paragraph()
        history_header = doc.add_paragraph("HISTORY OF PRESENT ILLNESS")
        if history_header.runs:
            history_header.runs[0].font.bold = True
            history_header.runs[0].font.size = Pt(14)
        
        history_text = getattr(template_data.medical_findings, 'history', None)
        if history_text:
            doc.add_paragraph(history_text)
        else:
            doc.add_paragraph("[HISTORY SECTION REQUIRES COMPLETION]")
        
        # Physical examination section
        doc.add_paragraph()
        exam_header = doc.add_paragraph("PHYSICAL EXAMINATION")
        if exam_header.runs:
            exam_header.runs[0].font.bold = True
            exam_header.runs[0].font.size = Pt(14)
        
        if template_data.medical_findings.findings:
            for finding in template_data.medical_findings.findings:
                doc.add_paragraph(f"• {finding}")
        else:
            doc.add_paragraph("[EXAMINATION FINDINGS REQUIRE COMPLETION]")
        
        # Diagnosis section
        doc.add_paragraph()
        diagnosis_header = doc.add_paragraph("DIAGNOSIS")
        if diagnosis_header.runs:
            diagnosis_header.runs[0].font.bold = True
            diagnosis_header.runs[0].font.size = Pt(14)
        
        if template_data.medical_findings.diagnoses:
            for i, diagnosis in enumerate(template_data.medical_findings.diagnoses, 1):
                doc.add_paragraph(f"{i}. {diagnosis.description}")
        else:
            doc.add_paragraph("[DIAGNOSES REQUIRE COMPLETION]")
        
        # Impairment rating section
        doc.add_paragraph()
        impairment_header = doc.add_paragraph("IMPAIRMENT RATING")
        if impairment_header.runs:
            impairment_header.runs[0].font.bold = True
            impairment_header.runs[0].font.size = Pt(14)
        
        if template_data.medical_findings.impairment_ratings:
            for rating in template_data.medical_findings.impairment_ratings:
                doc.add_paragraph(f"• {rating.percentage}% whole person impairment")
                body_part = getattr(rating, 'body_part', None)
                if body_part:
                    doc.add_paragraph(f"  Body part: {body_part}")
                if rating.ama_table:
                    doc.add_paragraph(f"  AMA Table: {rating.ama_table}")
        else:
            doc.add_paragraph("[IMPAIRMENT RATING REQUIRES COMPLETION]")
        
        # Conclusions section
        doc.add_paragraph()
        conclusion_header = doc.add_paragraph("CONCLUSIONS AND RECOMMENDATIONS")
        if conclusion_header.runs:
            conclusion_header.runs[0].font.bold = True
            conclusion_header.runs[0].font.size = Pt(14)
        
        doc.add_paragraph("[CONCLUSIONS AND RECOMMENDATIONS REQUIRE COMPLETION]")
    
    def _generate_output_path(self, template_data: QMETemplateData) -> str:
        """Generate output file path based on template data."""
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        patient_name = template_data.patient_info.name or "Unknown"
        
        # Sanitize patient name for filename
        safe_name = "".join(c for c in patient_name if c.isalnum() or c in (' ', '-', '_')).strip()
        safe_name = safe_name.replace(' ', '_')
        
        # Create results directory if it doesn't exist
        results_dir = Path("results/generated_documents")
        results_dir.mkdir(parents=True, exist_ok=True)
        
        filename = f"QME_Report_{safe_name}_{timestamp}.{self.template_config.output_format}"
        return str(results_dir / filename)
    
    def _get_file_statistics(self, file_path: str) -> Dict[str, int]:
        """Get file statistics for the generated document."""
        try:
            stats = {'size': 0, 'pages': 1, 'words': 0}
            
            if os.path.exists(file_path):
                stats['size'] = os.path.getsize(file_path)
                
                if file_path.endswith('.docx'):
                    try:
                        doc = Document(file_path)
                        # Estimate pages based on paragraph count
                        stats['pages'] = max(1, len(doc.paragraphs) // 25)
                        
                        # Count words
                        word_count = 0
                        for para in doc.paragraphs:
                            word_count += len(para.text.split())
                        stats['words'] = word_count
                    except:
                        pass  # Use defaults if document can't be read
            
            return stats
            
        except Exception as e:
            logger.error(f"Error getting file statistics: {e}")
            return {'size': 0, 'pages': 1, 'words': 0}
    
    def _attempt_error_recovery(self, template_data: QMETemplateData, original_error: Exception) -> AssemblyResult:
        """Attempt error recovery using fallback strategies."""
        logger.info("Attempting error recovery")
        
        for strategy in self.template_config.fallback_strategies:
            if self.error_recovery.recovery_attempts >= self.template_config.max_retry_attempts:
                break
                
            try:
                result = self.error_recovery.attempt_recovery(template_data, original_error, strategy)
                if result.status == AssemblyStatus.RECOVERED:
                    logger.info(f"Recovery successful using strategy: {strategy.value}")
                    return result
            except Exception as e:
                logger.error(f"Recovery strategy {strategy.value} failed: {e}")
                continue
        
        # All recovery attempts failed
        return AssemblyResult(
            status=AssemblyStatus.FAILED,
            template_data=template_data,
            error_message=f"All recovery attempts failed. Original error: {str(original_error)}",
            recovery_attempts=self.error_recovery.recovery_attempts
        )
    
    def validate_assembly_quality(self, result: AssemblyResult) -> QualityReport:
        """Validate the quality of an assembled template."""
        if not result.file_path or not os.path.exists(result.file_path):
            return QualityReport(
                overall_score=0.0,
                completeness_score=0.0,
                accuracy_score=0.0,
                compliance_score=0.0,
                formatting_score=0.0,
                total_issues=1,
                critical_issues=1,
                high_issues=0,
                medium_issues=0,
                low_issues=0,
                missing_sections=[],
                placeholder_count=0,
                compliance_status="error",
                recommendations=["File not found or inaccessible"]
            )
        
        return self.validator.validate_post_assembly(result.file_path, result.template_data)
    
    def apply_professional_formatting(self, document_path: str) -> bool:
        """Apply professional formatting to an existing document."""
        try:
            if not os.path.exists(document_path):
                logger.error(f"Document not found: {document_path}")
                return False
            
            doc = Document(document_path)
            formatted_doc = self.formatter.apply_professional_formatting(doc)
            formatted_doc.save(document_path)
            
            logger.info(f"Professional formatting applied to: {document_path}")
            return True
            
        except Exception as e:
            logger.error(f"Error applying professional formatting: {e}")
            return False