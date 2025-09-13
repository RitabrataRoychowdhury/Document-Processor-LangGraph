"""
Factory for creating document processors based on file type.
Implements the Factory Pattern for document processing.
"""

from abc import ABC, abstractmethod
from typing import Dict, Type, Optional, Tuple
import os
import io
try:
    from src.utils.logging_config import get_logger
except ImportError:
    from utils.logging_config import get_logger

# Conditional imports for external dependencies
try:
    import PyPDF2
    PYPDF2_AVAILABLE = True
except ImportError:
    PYPDF2_AVAILABLE = False

try:
    from docx import Document as DocxDocument
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False

logger = get_logger(__name__)


class DocumentProcessor(ABC):
    """Abstract base class for document processors."""
    
    @abstractmethod
    def extract_text(self, file_path_or_stream) -> str:
        """
        Extract text from document.
        
        Args:
            file_path_or_stream: File path string or file-like object
            
        Returns:
            Extracted text content
        """
        pass
    
    @abstractmethod
    def extract_metadata(self, file_path_or_stream) -> Dict[str, any]:
        """
        Extract metadata from document.
        
        Args:
            file_path_or_stream: File path string or file-like object
            
        Returns:
            Dictionary containing document metadata
        """
        pass
    
    @abstractmethod
    def get_supported_extensions(self) -> list[str]:
        """
        Get list of supported file extensions.
        
        Returns:
            List of supported file extensions (e.g., ['.pdf'])
        """
        pass


class PDFProcessor(DocumentProcessor):
    """Processor for PDF documents."""
    
    def extract_text(self, file_path_or_stream) -> str:
        """Extract text from PDF file."""
        if not PYPDF2_AVAILABLE:
            raise Exception("PyPDF2 is not installed. Please install it to process PDF files.")
        
        try:
            # Handle both file paths and file-like objects
            if isinstance(file_path_or_stream, str):
                with open(file_path_or_stream, 'rb') as file:
                    pdf_reader = PyPDF2.PdfReader(file)
                    return self._extract_pdf_text(pdf_reader)
            else:
                # File-like object (e.g., Streamlit UploadedFile)
                file_path_or_stream.seek(0)
                pdf_reader = PyPDF2.PdfReader(file_path_or_stream)
                return self._extract_pdf_text(pdf_reader)
                
        except Exception as e:
            logger.error(f"PDF text extraction error: {str(e)}")
            raise Exception(f"Failed to extract text from PDF: {str(e)}")
    
    def extract_metadata(self, file_path_or_stream) -> Dict[str, any]:
        """Extract metadata from PDF file."""
        if not PYPDF2_AVAILABLE:
            return {"error": "PyPDF2 is not installed"}
        
        try:
            if isinstance(file_path_or_stream, str):
                with open(file_path_or_stream, 'rb') as file:
                    pdf_reader = PyPDF2.PdfReader(file)
                    return self._extract_pdf_metadata(pdf_reader)
            else:
                file_path_or_stream.seek(0)
                pdf_reader = PyPDF2.PdfReader(file_path_or_stream)
                return self._extract_pdf_metadata(pdf_reader)
                
        except Exception as e:
            logger.error(f"PDF metadata extraction error: {str(e)}")
            return {"error": str(e)}
    
    def get_supported_extensions(self) -> list[str]:
        """Get supported file extensions for PDF processor."""
        return ['.pdf']
    
    def _extract_pdf_text(self, pdf_reader) -> str:
        """Extract text from PDF reader object."""
        text_content = []
        
        for page_num, page in enumerate(pdf_reader.pages):
            try:
                page_text = page.extract_text()
                if page_text.strip():
                    text_content.append(page_text)
            except Exception as e:
                logger.warning(f"Could not extract text from page {page_num + 1}: {str(e)}")
                continue
        
        if not text_content:
            raise Exception("No readable text found in PDF file")
        
        return "\n\n".join(text_content)
    
    def _extract_pdf_metadata(self, pdf_reader) -> Dict[str, any]:
        """Extract metadata from PDF reader object."""
        metadata = {
            "page_count": len(pdf_reader.pages),
            "processor_type": "PDF"
        }
        
        # Try to extract PDF metadata
        try:
            if pdf_reader.metadata:
                pdf_meta = pdf_reader.metadata
                metadata.update({
                    "title": pdf_meta.get("/Title", ""),
                    "author": pdf_meta.get("/Author", ""),
                    "subject": pdf_meta.get("/Subject", ""),
                    "creator": pdf_meta.get("/Creator", ""),
                    "producer": pdf_meta.get("/Producer", ""),
                    "creation_date": str(pdf_meta.get("/CreationDate", "")),
                    "modification_date": str(pdf_meta.get("/ModDate", ""))
                })
        except Exception as e:
            logger.warning(f"Could not extract PDF metadata: {e}")
            metadata["metadata_error"] = str(e)
        
        return metadata


class DOCXProcessor(DocumentProcessor):
    """Processor for DOCX documents."""
    
    def extract_text(self, file_path_or_stream) -> str:
        """Extract text from DOCX file."""
        if not DOCX_AVAILABLE:
            raise Exception("python-docx is not installed. Please install it to process DOCX files.")
        
        try:
            if isinstance(file_path_or_stream, str):
                doc = DocxDocument(file_path_or_stream)
            else:
                file_path_or_stream.seek(0)
                doc = DocxDocument(file_path_or_stream)
            
            text_content = []
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    text_content.append(paragraph.text)
            
            if not text_content:
                raise Exception("No readable text found in DOCX file")
            
            return "\n\n".join(text_content)
            
        except Exception as e:
            logger.error(f"DOCX text extraction error: {str(e)}")
            raise Exception(f"Failed to extract text from DOCX: {str(e)}")
    
    def extract_metadata(self, file_path_or_stream) -> Dict[str, any]:
        """Extract metadata from DOCX file."""
        if not DOCX_AVAILABLE:
            return {"error": "python-docx is not installed"}
        
        try:
            if isinstance(file_path_or_stream, str):
                doc = DocxDocument(file_path_or_stream)
            else:
                file_path_or_stream.seek(0)
                doc = DocxDocument(file_path_or_stream)
            
            metadata = {
                "paragraph_count": len(doc.paragraphs),
                "processor_type": "DOCX"
            }
            
            # Try to extract document properties
            try:
                core_props = doc.core_properties
                metadata.update({
                    "title": core_props.title or "",
                    "author": core_props.author or "",
                    "subject": core_props.subject or "",
                    "keywords": core_props.keywords or "",
                    "comments": core_props.comments or "",
                    "created": str(core_props.created) if core_props.created else "",
                    "modified": str(core_props.modified) if core_props.modified else "",
                    "last_modified_by": core_props.last_modified_by or ""
                })
            except Exception as e:
                logger.warning(f"Could not extract DOCX properties: {e}")
                metadata["metadata_error"] = str(e)
            
            return metadata
            
        except Exception as e:
            logger.error(f"DOCX metadata extraction error: {str(e)}")
            return {"error": str(e)}
    
    def get_supported_extensions(self) -> list[str]:
        """Get supported file extensions for DOCX processor."""
        return ['.docx']


class TXTProcessor(DocumentProcessor):
    """Processor for plain text documents."""
    
    def extract_text(self, file_path_or_stream) -> str:
        """Extract text from TXT file."""
        try:
            if isinstance(file_path_or_stream, str):
                # Try different encodings for file path
                encodings = ['utf-8', 'utf-16', 'latin-1', 'cp1252']
                for encoding in encodings:
                    try:
                        with open(file_path_or_stream, 'r', encoding=encoding) as file:
                            content = file.read()
                            if content.strip():
                                return content
                    except UnicodeDecodeError:
                        continue
                raise Exception("Could not decode text file with any supported encoding")
            else:
                # File-like object
                file_path_or_stream.seek(0)
                encodings = ['utf-8', 'utf-16', 'latin-1', 'cp1252']
                
                for encoding in encodings:
                    try:
                        file_path_or_stream.seek(0)
                        content = file_path_or_stream.read()
                        if isinstance(content, bytes):
                            text = content.decode(encoding)
                        else:
                            text = str(content)
                        
                        if text.strip():
                            return text
                    except UnicodeDecodeError:
                        continue
                
                raise Exception("Could not decode text file with any supported encoding")
                
        except Exception as e:
            logger.error(f"TXT text extraction error: {str(e)}")
            raise Exception(f"Failed to extract text from TXT: {str(e)}")
    
    def extract_metadata(self, file_path_or_stream) -> Dict[str, any]:
        """Extract metadata from TXT file."""
        try:
            text = self.extract_text(file_path_or_stream)
            
            # Basic text statistics
            lines = text.split('\n')
            words = text.split()
            
            metadata = {
                "line_count": len(lines),
                "word_count": len(words),
                "character_count": len(text),
                "processor_type": "TXT"
            }
            
            # Try to get file stats if it's a file path
            if isinstance(file_path_or_stream, str):
                try:
                    stat = os.stat(file_path_or_stream)
                    metadata.update({
                        "file_size": stat.st_size,
                        "created": str(stat.st_ctime),
                        "modified": str(stat.st_mtime)
                    })
                except Exception as e:
                    logger.warning(f"Could not get file stats: {e}")
            
            return metadata
            
        except Exception as e:
            logger.error(f"TXT metadata extraction error: {str(e)}")
            return {"error": str(e)}
    
    def get_supported_extensions(self) -> list[str]:
        """Get supported file extensions for TXT processor."""
        return ['.txt']


class ProcessorFactory:
    """Factory for creating document processors based on file type."""
    
    def __init__(self):
        """Initialize the processor factory with default processors."""
        self._processors: Dict[str, Type[DocumentProcessor]] = {}
        self._register_default_processors()
    
    def _register_default_processors(self):
        """Register default document processors based on available dependencies."""
        if PYPDF2_AVAILABLE:
            self.register_processor('.pdf', PDFProcessor)
        else:
            logger.warning("PyPDF2 not available - PDF processing disabled")
            
        if DOCX_AVAILABLE:
            self.register_processor('.docx', DOCXProcessor)
        else:
            logger.warning("python-docx not available - DOCX processing disabled")
            
        # TXT processor has no external dependencies
        self.register_processor('.txt', TXTProcessor)
    
    def register_processor(self, file_extension: str, processor_class: Type[DocumentProcessor]):
        """
        Register a processor for a specific file extension.
        
        Args:
            file_extension: File extension (e.g., '.pdf')
            processor_class: Processor class to handle this file type
        """
        if not file_extension.startswith('.'):
            file_extension = '.' + file_extension
        
        file_extension = file_extension.lower()
        self._processors[file_extension] = processor_class
        logger.info(f"Registered processor {processor_class.__name__} for {file_extension}")
    
    def create_processor(self, file_type: str) -> DocumentProcessor:
        """
        Create a processor for the specified file type.
        
        Args:
            file_type: File extension (e.g., '.pdf' or 'pdf')
            
        Returns:
            DocumentProcessor instance for the file type
            
        Raises:
            ValueError: If no processor is registered for the file type
        """
        # Normalize file type
        if not file_type.startswith('.'):
            file_type = '.' + file_type
        file_type = file_type.lower()
        
        if file_type not in self._processors:
            supported_types = list(self._processors.keys())
            raise ValueError(f"No processor registered for file type '{file_type}'. Supported types: {supported_types}")
        
        processor_class = self._processors[file_type]
        return processor_class()
    
    def get_supported_types(self) -> list[str]:
        """
        Get list of all supported file types.
        
        Returns:
            List of supported file extensions
        """
        return list(self._processors.keys())
    
    def is_supported(self, file_type: str) -> bool:
        """
        Check if a file type is supported.
        
        Args:
            file_type: File extension to check
            
        Returns:
            True if supported, False otherwise
        """
        if not file_type.startswith('.'):
            file_type = '.' + file_type
        file_type = file_type.lower()
        
        return file_type in self._processors


# Global factory instance
processor_factory = ProcessorFactory()